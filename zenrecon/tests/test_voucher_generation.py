# -*- coding: utf-8 -*-
import base64
import io

from openpyxl import Workbook, load_workbook

from odoo.exceptions import UserError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged("post_install", "-at_install")
class TestMultiReportVoucherConfiguration(TransactionCase):
    def _selected_reports(self, *codes):
        options = self.env["bsr.accounting.report.code"].search([("code", "in", list(codes))])
        return [(6, 0, options.ids)]

    def test_mixed_currency_values_share_one_row_and_balance_at_product_level(self):
        run = self.env["bsr.accounting.run"].new(
            {
                "product_type": "Prepaid",
                "usd_buy_rate": 123.0,
                "usd_sale_rate": 122.0,
                "usd_mid_rate": 120.0,
                "currency_conversion_gl_number": "5555555555555",
            }
        )

        def group(report_code, amount_source, amount, side, gl_number, clearing=False, extra=None):
            configuration = self.env["bsr.accounting.configuration"].new(
                {
                    "business_owner": "MGBL",
                    "product_type": "Prepaid",
                    "reporting_for": "1000675529 PREPAID 407275",
                    "settlement_currency": "USD",
                    "clearing_currency": clearing,
                }
            )
            mapping = self.env["bsr.accounting.mapping"].new(
                {
                    "configuration_id": configuration,
                    "report_code": report_code,
                    "amount_source": amount_source,
                    "matched_amount": amount,
                    "posting_indicator": side,
                    "source_currency": clearing or "USD",
                    "clearing_currency": clearing or "",
                    "conversion_type": "mid",
                    "gl_number": gl_number,
                    "gl_name": gl_number,
                    "amount_columns_json": extra or {},
                }
            )
            report = self.env["bsr.accounting.report"].new(
                {"configuration_id": configuration}
            )
            return report, mapping

        groups = [
            group("VSS-110", "total_amount", 0.17, "CR", "00000"),
            group("VSS-110", "total_amount", 14.09, "CR", "1111"),
            group("VSS-110", "total_amount", 0.14, "DR", "33333"),
            group(
                "VSS-120", "clearing_amount", 1778.84, "DR", "1313",
                clearing="BDT", extra={"interchange_debit": 14.12},
            ),
            group(
                "VSS-210", "optional_issuer_fee", 34.88, "CR", "09099",
                clearing="BDT",
            ),
        ]
        data = run._voucher_section_data(groups)

        self.assertAlmostEqual(data["debit_usd"], 14.26, places=2)
        self.assertAlmostEqual(data["credit_usd"], 14.26, places=2)
        self.assertAlmostEqual(data["debit_bdt"], 1795.64, places=2)
        self.assertAlmostEqual(data["credit_bdt"], 1795.64, places=2)
        self.assertEqual(data["adjustment"]["side"], "CR")
        self.assertAlmostEqual(data["adjustment"]["amount_bdt"], 49.56, places=2)
        self.assertFalse(data["additional_fields"])
        values_by_gl = {
            item["mapping"].gl_number: (item["amount_usd"], item["amount_bdt"])
            for item in data["lines"]
        }
        self.assertEqual(values_by_gl["1313"], (14.12, 1778.84))
        self.assertEqual(values_by_gl["09099"], (0.0, 34.88))

    def test_bdt_adjustment_does_not_hide_usd_mismatch(self):
        run = self.env["bsr.accounting.run"].new(
            {
                "usd_mid_rate": 123.0,
                "currency_conversion_gl_number": "777777777777",
            }
        )
        configuration = self.env["bsr.accounting.configuration"].new(
            {
                "settlement_currency": "USD",
                # This reusable filter is stale; the runtime mapping below is
                # authoritative for placement in the voucher columns.
                "clearing_currency": "USD",
            }
        )
        debit = self.env["bsr.accounting.mapping"].new(
            {
                "configuration_id": configuration,
                "report_code": "VSS-120",
                "amount_source": "clearing_amount",
                "matched_amount": 107268.32,
                "posting_indicator": "DR",
                "source_currency": "BDT",
                "clearing_currency": "BDT",
                "gl_number": "1313",
                "gl_name": "ATM USD BDT",
                "debit_remarks": "BSR/CONFIG/2026/00153 — FT SRE BDT ATM reconciliation posting",
                "amount_columns_json": {"interchange_debit": 855.0},
            }
        )
        credit = self.env["bsr.accounting.mapping"].new(
            {
                "configuration_id": configuration,
                "report_code": "VSS-110",
                "amount_source": "total_amount",
                "matched_amount": 800.0,
                "posting_indicator": "CR",
                "source_currency": "USD",
                "conversion_type": "mid",
                "gl_number": "1111",
                "gl_name": "Settlement",
            }
        )
        report = self.env["bsr.accounting.report"].new(
            {"configuration_id": configuration}
        )

        data = run._voucher_section_data([(report, debit | credit)])

        source_line = next(
            line for line in data["lines"] if line["mapping"].gl_number == "1313"
        )
        self.assertEqual(source_line["amount_usd"], 855.0)
        self.assertEqual(source_line["amount_bdt"], 107268.32)
        self.assertFalse(source_line["additional"])
        self.assertAlmostEqual(data["adjustment"]["amount_usd"], 0.0, places=2)
        self.assertAlmostEqual(data["adjustment"]["amount_bdt"], 8868.32, places=2)
        self.assertAlmostEqual(data["debit_usd"], 855.0, places=2)
        self.assertAlmostEqual(data["credit_usd"], 800.0, places=2)
        self.assertAlmostEqual(data["debit_bdt"], data["credit_bdt"], places=2)
        self.assertAlmostEqual(data["usd_difference"], 55.0, places=2)
        self.assertAlmostEqual(data["bdt_difference"], 0.0, places=2)
        self.assertAlmostEqual(data["difference"], 55.0, places=2)

        workbook = Workbook()
        sheet = workbook.active
        _last_row, header_row, table_end = run._write_voucher_table(sheet, 1, data)
        self.assertEqual(table_end, 10)
        self.assertEqual(
            [sheet.cell(header_row, column).value for column in range(1, 11)],
            [
                "GL Number", "GL Name", "Debit Amount (USD)", "Debit Amount (BDT)",
                "Debit Remarks", "GL Number", "GL Name", "Credit Amount (USD)",
                "Credit Amount (BDT)", "Credit Remarks",
            ],
        )
        self.assertNotIn("Interchange Debit", [cell.value for cell in sheet[header_row]])
        source_row = next(
            row for row in range(header_row + 1, sheet.max_row + 1)
            if sheet.cell(row, 1).value == "1313"
        )
        self.assertTrue(sheet.cell(source_row, 5).alignment.wrap_text)
        self.assertEqual(sheet.cell(source_row, 5).alignment.vertical, "center")
        self.assertGreater(sheet.row_dimensions[source_row].height, 24)

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        attachment = cls.env["ir.attachment"].create(
            {"name": "EP747.txt", "raw": b"test source", "mimetype": "text/plain"}
        )
        cls.vss_import = cls.env["bsr.vss.import"].create(
            {
                "attachment_id": attachment.id,
                "company_id": cls.env.company.id,
                "state": "completed",
                "file_name": "EP747.txt",
            }
        )
        common = {
            "import_id": cls.vss_import.id,
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "1000512940 BIN 474867 NNSS",
            "rollup_to_display": "1000137367 THE CITY DOM",
            "funds_xfer_entity_display": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
        }
        cls.env["bsr.vss.120.line"].create(
            [
                dict(
                    common,
                    report_code="VSS-120",
                    clearing_currency="BDT",
                    business_mode="ISSUER",
                    transaction_category="PURCHASE",
                    transaction_type="ORIGINAL SALE",
                    disposition="",
                    clearing_direction="DB",
                    clearing_amount=946810.59,
                    record_type="detail",
                ),
                dict(
                    common,
                    report_code="VSS-120",
                    clearing_currency="BDT",
                    business_mode="ISSUER",
                    transaction_category="PURCHASE",
                    transaction_type="TOTAL PURCHASE",
                    disposition="",
                    clearing_direction="DB",
                    clearing_amount=946810.59,
                    record_type="total",
                ),
                dict(
                    common,
                    report_code="VSS-120",
                    clearing_currency="BDT",
                    business_mode="ISSUER",
                    transaction_category="MERCHANDISE CREDIT",
                    transaction_type="ORIGINAL",
                    disposition="",
                    clearing_direction="CR",
                    clearing_amount=399.00,
                    record_type="detail",
                ),
                dict(
                    common,
                    report_code="VSS-120",
                    clearing_currency="BDT",
                    business_mode="ISSUER",
                    transaction_category="MERCHANDISE CREDIT",
                    transaction_type="TOTAL MERCHANDISE CREDIT",
                    disposition="",
                    clearing_direction="CR",
                    clearing_amount=399.00,
                    record_type="total",
                ),
            ]
        )
        cls.env["bsr.vss.130.line"].create(
            [
                dict(
                    common,
                    report_code="VSS-130",
                    business_mode="ISSUER",
                    transaction_category="PURCHASE",
                    transaction_type="ORIGINAL SALE",
                    description="AUTHENT SEC PLAT",
                    interchange_direction="DB",
                    interchange_amount=946810.59,
                    fee_credit=16528.35,
                    record_type="detail",
                ),
                dict(
                    common,
                    report_code="VSS-130",
                    business_mode="ISSUER",
                    transaction_category="MERCHANDISE CREDIT",
                    transaction_type="ORIGINAL",
                    description="TOTAL ISSUER REIMB FEES",
                    fee_credit=16528.35,
                    fee_debit=7.38,
                    record_type="total",
                ),
                dict(
                    common,
                    report_code="VSS-130",
                    business_mode="ISSUER",
                    transaction_category="MERCHANDISE CREDIT",
                    transaction_type="ORIGINAL",
                    description="NET ISSUER REIMB FEES",
                    fee_credit=16520.97,
                    record_type="net",
                ),
            ]
        )
        cls.env["bsr.vss.110.line"].create(
            [
                dict(
                    common,
                    report_code="VSS-110",
                    particulars="INTERCHANGE VALUE",
                    entity="TOTAL ISSUER",
                    credit_amount=399.00,
                    debit_amount=946810.59,
                    total_amount=946411.59,
                    direction="DB",
                    record_type="detail",
                ),
                dict(
                    common,
                    report_code="VSS-110",
                    particulars="INTERCHANGE VALUE",
                    entity="TOTAL INTERCHANGE VALUE",
                    credit_amount=399.00,
                    debit_amount=946810.59,
                    total_amount=946411.59,
                    direction="DB",
                    record_type="total",
                ),
                dict(
                    common,
                    report_code="VSS-110",
                    particulars="REIMBURSEMENT FEES",
                    entity="TOTAL REIMBURSEMENT FEES",
                    credit_amount=16528.35,
                    debit_amount=7.38,
                    total_amount=16520.97,
                    direction="CR",
                    record_type="total",
                ),
                dict(
                    common,
                    report_code="VSS-110",
                    particulars="TOTAL",
                    entity="NET SETTLEMENT AMOUNT",
                    total_amount=929890.62,
                    direction="DB",
                    record_type="net",
                ),
            ]
        )
        acquirer_common = {
            "import_id": cls.vss_import.id,
            "report_code": "VSS-130",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "1000548730 406591 SMS ACQ",
            "rollup_to_display": "1000137367 THE CITY DOM",
            "funds_xfer_entity_display": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
            "business_mode": "ACQUIRER",
            "transaction_category": "ORIGINAL CREDIT",
            "transaction_type": "ORIGINAL",
        }
        cls.env["bsr.vss.130.line"].create(
            [
                dict(
                    acquirer_common,
                    description="ORIGINAL CREDIT",
                    interchange_direction="DB",
                    interchange_amount=755191.65,
                    fee_credit=0.0,
                    fee_debit=416.0,
                    record_type="detail",
                ),
                dict(
                    acquirer_common,
                    description="TOTAL ACQUIRER REIMB FEES",
                    fee_credit=0.0,
                    fee_debit=416.0,
                    record_type="total",
                ),
            ]
        )

    def test_visa_setup_source_choices_and_configuration_grouping(self):
        configuration_values = {
            "source_import_id": self.vss_import.id,
            "business_owner": "Retail",
            "product_type": "Credit Card",
            "reporting_for": "1000512940 BIN 474867 NNSS",
            "rollup_to": "1000137367 THE CITY DOM",
            "funds_xfer_entity": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
        }
        configurations = self.env["bsr.accounting.configuration"].create(
            [configuration_values, dict(configuration_values)]
        )
        options = self.env["bsr.visa.source.option"].search(
            [("company_id", "=", self.env.company.id)]
        )
        reporting_options = options.filtered(
            lambda option: option.field_name == "reporting_for"
            and option.name == "1000512940 BIN 474867 NNSS"
        )
        rollup_options = options.filtered(
            lambda option: option.field_name == "rollup_to"
            and option.name == "1000137367 THE CITY DOM"
        )
        funds_options = options.filtered(
            lambda option: option.field_name == "funds_xfer_entity"
            and option.name == "1000137367 THE CITY DOM"
        )
        self.assertEqual(
            set(reporting_options.mapped("name")),
            {"1000512940 BIN 474867 NNSS"},
        )
        self.assertEqual(rollup_options.mapped("name"), ["1000137367 THE CITY DOM"])
        self.assertEqual(funds_options.mapped("name"), ["1000137367 THE CITY DOM"])

        setup = self.env["bsr.visa.setup"].create(
            {
                "reporting_for_option_ids": [(6, 0, reporting_options.ids)],
                "rollup_to_option_ids": [(6, 0, rollup_options.ids)],
                "funds_xfer_entity_option_ids": [(6, 0, funds_options.ids)],
                "business_owner": "Retail",
                "product_type": "Credit Card",
                "card_type": "Platinum",
            }
        )
        self.assertEqual(len(setup.reporting_for_option_ids), 1)
        self.assertEqual(setup.card_type, "Platinum")
        dimension_values = self.env["bsr.visa.dimension.option"].search(
            [("company_id", "=", self.env.company.id)]
        )
        self.assertIn("Retail", dimension_values.mapped("name"))
        self.assertIn("All", dimension_values.mapped("name"))

        wizard = self.env["bsr.configuration.group.wizard"].create(
            {
                "configuration_ids": [(6, 0, configurations.ids)],
                "company_id": self.env.company.id,
                "new_group_name": "USD Voucher",
                "settlement_currency": "usd",
            }
        )
        wizard.action_apply()
        self.assertEqual(
            set(configurations.mapped("configuration_group_id.name")),
            {"USD Voucher"},
        )
        self.assertEqual(configurations.configuration_group_id.configuration_count, 2)
        self.assertEqual(configurations.configuration_group_id.settlement_currency, "usd")

        self.env["bsr.configuration.group.wizard"].create(
            {
                "operation": "remove",
                "configuration_ids": [(6, 0, configurations[:1].ids)],
                "company_id": self.env.company.id,
            }
        ).action_apply()
        self.assertFalse(configurations[:1].configuration_group_id)
        self.assertEqual(configurations[1:].configuration_group_id.name, "USD Voucher")

    def test_run_scope_uses_all_groups_with_the_selected_currency(self):
        configuration_values = {
            "source_import_id": self.vss_import.id,
            "business_owner": "Retail",
            "product_type": "Credit Card",
            "reporting_for": "1000512940 BIN 474867 NNSS",
            "rollup_to": "1000137367 THE CITY DOM",
            "funds_xfer_entity": "1000137367 THE CITY DOM",
            "settlement_currency": "USD",
            "state": "ready",
        }
        configurations = self.env["bsr.accounting.configuration"].create(
            [configuration_values, dict(configuration_values)]
        )
        bdt_configuration = self.env["bsr.accounting.configuration"].create(
            dict(configuration_values, settlement_currency="BDT")
        )
        groups = self.env["bsr.accounting.configuration.group"].create(
            [
                {"name": "USD Voucher A", "settlement_currency": "usd"},
                {"name": "USD Voucher B", "settlement_currency": "usd"},
            ]
        )
        all_currency_group = self.env["bsr.accounting.configuration.group"].create(
            {"name": "All Currency Voucher", "settlement_currency": "all"}
        )
        configurations[0].configuration_group_id = groups[0]
        configurations[1].configuration_group_id = groups[1]
        bdt_configuration.configuration_group_id = all_currency_group

        options = self.env["bsr.visa.source.option"].search(
            [("company_id", "=", self.env.company.id)]
        )
        self.env["bsr.visa.setup"].create(
            {
                "reporting_for_option_ids": [
                    (6, 0, options.filtered(
                        lambda row: row.field_name == "reporting_for"
                        and row.name == "1000512940 BIN 474867 NNSS"
                    ).ids)
                ],
                "rollup_to_option_ids": [
                    (6, 0, options.filtered(
                        lambda row: row.field_name == "rollup_to"
                        and row.name == "1000137367 THE CITY DOM"
                    ).ids)
                ],
                "funds_xfer_entity_option_ids": [
                    (6, 0, options.filtered(
                        lambda row: row.field_name == "funds_xfer_entity"
                        and row.name == "1000137367 THE CITY DOM"
                    ).ids)
                ],
                "business_owner": "Retail",
                "product_type": "Credit Card",
                "card_type": "Platinum",
            }
        )
        Dimension = self.env["bsr.visa.dimension.option"]
        run = self.env["bsr.accounting.run"].create(
            {
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "business_owner_filter_id": Dimension.search(
                    [("dimension", "=", "business_owner"), ("name", "=", "Retail")], limit=1
                ).id,
                "product_type_filter_id": Dimension.search(
                    [("dimension", "=", "product_type"), ("name", "=", "Credit Card")], limit=1
                ).id,
                "card_type_filter_id": Dimension.search(
                    [("dimension", "=", "card_type"), ("name", "=", "Platinum")], limit=1
                ).id,
                "settlement_currency_scope": "usd",
            }
        )
        run._prepare_visa_run_scope()

        self.assertEqual(run.source_import_id, self.vss_import)
        self.assertTrue(set(configurations.ids).issubset(set(run.configuration_ids.ids)))
        self.assertTrue(
            set(groups.ids).issubset(set(run.configuration_ids.configuration_group_id.ids))
        )
        self.assertNotIn(bdt_configuration, run.configuration_ids)
        self.assertNotIn(all_currency_group, run.configuration_ids.configuration_group_id)
        self.assertEqual(
            run.product_type_option_ids.mapped("name"),
            ["1000512940 BIN 474867 NNSS"],
        )

        explicit_mismatch_run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": self.vss_import.id,
                "configuration_ids": [(6, 0, bdt_configuration.ids)],
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "settlement_currency_scope": "usd",
            }
        )
        with self.assertRaisesRegex(UserError, "Settlement Currency mismatch"):
            explicit_mismatch_run._validate_selected_configuration_currencies()

    def test_client_platinum_voucher_is_balanced_without_duplicate_levels(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120", "VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Platinum Credit Card/Mother SRE",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "remarks": "Daily settlement voucher",
            }
        )
        configuration.action_match_source()

        self.assertEqual(configuration.state, "matched")
        self.assertEqual(len(configuration.line_ids), 4)
        self.assertEqual(configuration.source_report_codes, "VSS-110, VSS-120, VSS-130")
        self.assertFalse(configuration.line_ids.filtered(lambda line: line.report_code in {"VSS-140", "VSS-210"}))
        self.assertEqual(set(configuration.line_ids.mapped("record_type")), {"detail", "net"})
        settlement = configuration.line_ids.filtered(lambda line: line.report_code == "VSS-110")
        self.assertEqual(settlement.matched_amount, 929890.62)
        self.assertEqual(settlement.source_direction, "DB")
        self.assertEqual(settlement.posting_indicator, "CR")
        self.assertEqual(configuration.preview_debit_total, 946810.59)
        self.assertEqual(configuration.preview_credit_total, 946810.59)
        self.assertTrue(configuration.is_preview_balanced)

        for line in configuration.line_ids:
            if line.report_code == "VSS-110":
                gl_number, gl_name = "00000000", "NET SETTLEMENT AMOUNT"
            elif line.report_code == "VSS-130":
                gl_number, gl_name = "55055505", "VISA IRF Income"
            else:
                gl_number, gl_name = "99999999", "VISA outstanding GL"
            line.write({"gl_number": gl_number, "gl_name": gl_name, "remarks": "Configured by administrator"})

        configuration.action_validate()
        self.assertEqual(configuration.state, "ready")

        action = configuration.action_generate_report()
        report = self.env["bsr.accounting.report"].browse(action["res_id"])
        self.assertEqual(report.state, "balanced")
        self.assertEqual(report.debit_total, 946810.59)
        self.assertEqual(report.credit_total, 946810.59)
        self.assertEqual(len(report.voucher_id.line_ids), 4)
        self.assertTrue(report.output_filename.endswith(".xlsx"))

        workbook = load_workbook(io.BytesIO(base64.b64decode(report.output_file)), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Voucher Report", "Source Details"])
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertIn("Report ID", values)
        self.assertIn("VSS-110, VSS-120, VSS-130", values)
        self.assertIn("Platinum Credit Card/Mother SRE", values)
        self.assertIn("Daily settlement voucher", " ".join(str(value or "") for value in values))
        header_row = next(row[0].row for row in workbook.active.iter_rows() if row[0].value == "GL Number")
        self.assertEqual(workbook.active.cell(header_row, 4).value, "Debit Remarks")
        self.assertEqual(workbook.active.cell(header_row, 8).value, "Credit Remarks")
        self.assertEqual(workbook.active.max_column, 8)
        self.assertEqual(workbook.active.cell(header_row + 1, 8).value, "Configured by administrator")
        details = workbook["Source Details"]
        self.assertEqual(details.max_row, 8)
        self.assertEqual([details.cell(row, 1).value for row in range(5, 9)],
                         ["VSS-110", "VSS-120", "VSS-120", "VSS-130"])
        gl_column = next(
            column
            for column in range(10, details.max_column + 1)
            if details.cell(4, column).value == "GL Number"
        )
        source_amounts = [
            next(
                value
                for column in range(10, gl_column)
                if (value := details.cell(row, column).value)
            )
            for row in range(5, 9)
        ]
        self.assertEqual(source_amounts, [929890.62, 946810.59, 399.00, 16520.97])

    def test_mismatch_can_be_generated_but_not_marked_balanced(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120", "VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Platinum Credit Card",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
            }
        )
        configuration.action_match_source()
        configuration.line_ids.write({"gl_number": "99999999", "gl_name": "Configured GL"})
        configuration.line_ids.filtered(lambda line: line.report_code == "VSS-130").include_in_voucher = False

        configuration.action_validate()
        self.assertEqual(configuration.state, "ready")
        self.assertFalse(configuration.is_preview_balanced)

        action = configuration.action_generate_report()
        report = self.env["bsr.accounting.report"].browse(action["res_id"])
        self.assertEqual(report.state, "exception")
        self.assertAlmostEqual(report.difference, 16520.97, places=2)
        workbook = load_workbook(io.BytesIO(base64.b64decode(report.output_file)), data_only=True)
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertIn("MISMATCH FLAG", values)

    def test_same_gl_keeps_separate_source_lines_in_xlsx(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120", "VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Platinum Credit Card/Mother SRE",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
            }
        )
        configuration.action_match_source()
        configuration.line_ids.write({"gl_number": "99999999", "gl_name": "VISA outstanding GL"})
        configuration.action_validate()
        action = configuration.action_generate_report()
        report = self.env["bsr.accounting.report"].browse(action["res_id"])

        self.assertEqual(report.state, "balanced")
        self.assertEqual(len(report.voucher_id.line_ids), 4)
        self.assertEqual(
            {line.source_report_code for line in report.voucher_id.line_ids},
            {"VSS-110", "VSS-120", "VSS-130"},
        )
        workbook = load_workbook(io.BytesIO(base64.b64decode(report.output_file)), data_only=True)
        sheet = workbook.active
        header_row = next(row[0].row for row in sheet.iter_rows() if row[0].value == "GL Number")
        self.assertEqual(sheet.cell(header_row + 1, 3).value, 946810.59)
        self.assertEqual(
            [sheet.cell(header_row + index, 7).value for index in (1, 2, 3)],
            [929890.62, 399.00, 16520.97],
        )

    def test_manual_report_and_source_columns_filter_exact_rows(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Selected categories",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE, MERCHANDISE CREDIT",
                "transaction_type": "ORIGINAL SALE, ORIGINAL",
            }
        )
        configuration.action_match_source()

        self.assertEqual(configuration.source_report_codes, "VSS-120")
        self.assertEqual(len(configuration.line_ids), 2)
        self.assertEqual(
            set(configuration.line_ids.mapped("transaction_category")),
            {"PURCHASE", "MERCHANDISE CREDIT"},
        )

        configuration.action_reset_draft()
        configuration.write(
            {
                "report_code_ids": self._selected_reports("VSS-110"),
                # These columns do not exist in VSS-110 and therefore must not
                # exclude an otherwise matching VSS-110 source row.
                "clearing_currency": "USD",
                "business_mode": "ACQUIRER",
                "transaction_category": "ATM CASH",
                "transaction_type": "CHARGEBACK",
                "source_description": "RVRSL",
                "particulars": "TOTAL",
                "entity": "NET SETTLEMENT AMOUNT",
            }
        )
        configuration.action_match_source()
        self.assertEqual(configuration.source_report_codes, "VSS-110")
        self.assertEqual(len(configuration.line_ids), 1)
        self.assertEqual(configuration.line_ids.description_match, "NET SETTLEMENT AMOUNT")

    def test_disposition_source_tags_distinguish_blank_from_reversal(self):
        common = {
            "import_id": self.vss_import.id,
            "report_code": "VSS-120",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "1000512940 BIN 474867 NNSS",
            "rollup_to_display": "1000137367 THE CITY DOM",
            "funds_xfer_entity_display": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
            "clearing_currency": "BDT",
            "business_mode": "ISSUER",
            "transaction_category": "PURCHASE",
            "transaction_type": "ORIGINAL SALE",
            "record_type": "detail",
        }
        self.env["bsr.vss.120.line"].create(
            dict(common, disposition="RVRSL", clearing_direction="CR", clearing_amount=8493.00)
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Trust Bank",
                "product_type": "Interchange Value Report",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "vss_120_source_level": "detail",
                "amount_vss_120_interchange_credit": False,
                "amount_vss_120_interchange_debit": False,
            }
        )

        options = self.env["bsr.accounting.filter.option"].search(
            [
                ("source_import_id", "=", self.vss_import.id),
                ("field_name", "=", "disposition"),
            ]
        )
        blank = options.filtered("is_blank")
        reversal = options.filtered(lambda option: option.value == "RVRSL")
        self.assertTrue(blank)
        self.assertTrue(reversal)
        self.assertEqual(set(options.report_code_ids.mapped("code")), {"VSS-120"})

        configuration.disposition_option_ids = blank
        configuration.action_match_source()
        self.assertEqual(configuration.line_ids.mapped("matched_amount"), [946810.59])

        cases = (
            (reversal, [8493.00]),
            (blank | reversal, [946810.59, 8493.00]),
            (options.browse(), [946810.59, 8493.00]),
        )
        for selected, expected in cases:
            configuration.disposition_option_ids = selected
            configuration.action_match_source()
            self.assertEqual(sorted(configuration.line_ids.mapped("matched_amount")), sorted(expected))

    def test_configuration_rules_are_isolated_preserved_and_deletable(self):
        attachment = self.env["ir.attachment"].create(
            {"name": "EP747-second.txt", "raw": b"second source", "mimetype": "text/plain"}
        )
        second_import = self.env["bsr.vss.import"].create(
            {
                "attachment_id": attachment.id,
                "company_id": self.env.company.id,
                "state": "completed",
                "file_name": "EP747-second.txt",
            }
        )
        self.env["bsr.vss.120.line"].create(
            {
                "import_id": second_import.id,
                "report_code": "VSS-120",
                "processing_date": "2026-09-15",
                "report_date": "2026-09-14",
                "reporting_for_display": "1000646689 PLAT CREDIT/I",
                "rollup_to_display": "1000200636 TRUST BANK",
                "funds_xfer_entity_display": "1000200636 TRUST BANK",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "clearing_direction": "DB",
                "clearing_amount": 17595120.16,
                "record_type": "detail",
            }
        )

        def make_configuration(source_import, production_date, report_date, reporting_for, rollup_to):
            configuration = self.env["bsr.accounting.configuration"].create(
                {
                    "source_import_id": source_import.id,
                    "report_code_ids": self._selected_reports("VSS-120"),
                    "production_date": production_date,
                    "report_date": report_date,
                    "posting_date": production_date,
                    "business_owner": "Trust Bank",
                    "product_type": "Interchange Value Report",
                    "reporting_for": reporting_for,
                    "rollup_to": rollup_to,
                    "funds_xfer_entity": rollup_to,
                    "settlement_currency": "BDT",
                    "clearing_currency": "BDT",
                    "business_mode": "ISSUER",
                    "transaction_category": "PURCHASE",
                    "transaction_type": "ORIGINAL SALE",
                    "vss_120_source_level": "detail",
                    "amount_vss_120_interchange_credit": False,
                    "amount_vss_120_interchange_debit": False,
                }
            )
            configuration.action_match_source()
            return configuration

        first = make_configuration(
            self.vss_import,
            "2018-03-01",
            "2018-03-01",
            "1000512940 BIN 474867 NNSS",
            "1000137367 THE CITY DOM",
        )
        second = make_configuration(
            second_import,
            "2026-09-15",
            "2026-09-14",
            "1000646689 PLAT CREDIT/I",
            "1000200636 TRUST BANK",
        )

        self.assertEqual(second.previous_same_report_line_ids, second.line_ids)
        self.assertNotIn(first, second.previous_same_report_line_ids.configuration_id)
        self.assertEqual(set(second.line_ids.mapped("report_code")), {"VSS-120"})
        self.assertEqual(first.preview_debit_total, 946810.59)
        self.assertEqual(second.preview_debit_total, 17595120.16)

        first.line_ids.write({"gl_number": "120200", "gl_name": "VSS-120 Clearing"})
        self.assertEqual(first.line_ids.gl_number, "120200")
        self.assertEqual(first.line_ids.gl_name, "VSS-120 Clearing")

        second.line_ids.write({"gl_number": "120201", "gl_name": "VSS-120 Current"})
        first.action_validate()
        second.action_validate()
        expected_line_count = len(second.line_ids.filtered("include_in_voucher"))
        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": second_import.id,
                "production_date": "2026-09-15",
                "report_date": "2026-09-14",
                "posting_date": "2026-09-26",
                "product_type": "Combined VSS-120",
                "configuration_ids": [(6, 0, second.ids)],
            }
        )
        run.action_execute()
        workbook = load_workbook(io.BytesIO(base64.b64decode(run.output_file)), data_only=True)
        details = workbook["Merged Source Details"]
        self.assertEqual(details.max_row - 1, expected_line_count)
        self.assertEqual(
            {details.cell(row, 3).value for row in range(2, details.max_row + 1)},
            {"VSS-120"},
        )

        second.action_reset_draft()
        editable_rule = second.line_ids
        self.assertTrue(editable_rule)
        self.assertEqual(second.previous_same_report_line_ids, editable_rule)
        editable_rule.action_delete_rule()
        self.assertFalse(second.line_ids)
        self.assertFalse(second.previous_same_report_line_ids)

    def test_available_reports_come_from_selected_import(self):
        configuration = self.env["bsr.accounting.configuration"].new(
            {"source_import_id": self.vss_import.id}
        )
        self.assertEqual(
            set(configuration.available_report_code_ids.mapped("code")),
            {"VSS-110", "VSS-120", "VSS-130"},
        )
        configuration._onchange_source_scope()
        self.assertFalse(configuration.report_code_ids)
        self.assertEqual(configuration.production_date.isoformat(), "2018-03-01")
        self.assertEqual(configuration.report_date.isoformat(), "2018-03-01")

    def test_report_specific_amount_sections_and_optional_vss_120_level(self):
        configuration = self.env["bsr.accounting.configuration"].new(
            {"report_code_ids": self._selected_reports("VSS-110", "VSS-130")}
        )
        self.assertTrue(configuration.show_vss_110_amounts)
        self.assertFalse(configuration.show_vss_120_amounts)
        self.assertTrue(configuration.show_vss_130_amounts)
        self.assertFalse(configuration.vss_120_source_level)
        configuration.report_code_ids = self.env["bsr.accounting.report.code"].search(
            [("code", "=", "VSS-120")]
        )
        configuration._onchange_report_code_ids()
        self.assertEqual(configuration.vss_120_source_level, "detail")

        configuration.report_code_ids = self.env["bsr.accounting.report.code"].search(
            [("code", "=", "VSS-110")]
        )
        configuration._onchange_report_code_ids()
        self.assertFalse(configuration.vss_120_source_level)

    def test_usd_amount_column_is_only_used_for_mixed_bdt_usd(self):
        self.env["bsr.vss.130.line"].create(
            {
                "import_id": self.vss_import.id,
                "report_code": "VSS-130",
                "processing_date": "2018-03-01",
                "report_date": "2018-03-01",
                "reporting_for_display": "USD TEST",
                "rollup_to_display": "USD TEST",
                "funds_xfer_entity_display": "USD TEST",
                "settlement_currency": "USD",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "description": "USD INTERCHANGE",
                "interchange_amount": 123.45,
                "interchange_direction": "DB",
                "record_type": "detail",
            }
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "USD Test",
                "product_type": "USD Test",
                "reporting_for": "USD TEST",
                "rollup_to": "USD TEST",
                "funds_xfer_entity": "USD TEST",
                "settlement_currency": "USD",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "amount_vss_130_fee_credit": False,
                "amount_vss_130_fee_debit": False,
            }
        )
        self.assertFalse(configuration.show_usd_amount)
        configuration.action_match_source()
        self.assertEqual(configuration.line_ids.mapped("usd_amount"), [0.0])
        configuration.line_ids.write({"gl_number": "USD-GL", "gl_name": "USD Clearing"})
        configuration.action_validate()
        report = self.env["bsr.accounting.report"].create({"configuration_id": configuration.id})
        report.action_generate()
        self.assertEqual(report.voucher_id.line_ids.usd_amount, 0.0)
        workbook = load_workbook(io.BytesIO(base64.b64decode(report.output_file)), data_only=True)
        sheet = workbook["Voucher Report"]
        self.assertNotIn("USD Amount", [cell.value for row in sheet.iter_rows() for cell in row])
        mixed = self.env["bsr.accounting.configuration"].new(
            {"settlement_currency": "USD", "clearing_currency": "BDT"}
        )
        self.assertTrue(mixed.show_usd_amount)
        mixed.clearing_currency = "USD"
        self.assertFalse(mixed.show_usd_amount)

    def test_vss_130_region_and_fee_category_options_filter_detail_rows(self):
        common = {
            "import_id": self.vss_import.id,
            "report_code": "VSS-130",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "1000512940 BIN 474867 NNSS",
            "rollup_to_display": "1000137367 THE CITY DOM",
            "funds_xfer_entity_display": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
            "business_mode": "ACQUIRER",
            "transaction_category": "PURCHASE",
            "transaction_type": "ORIGINAL SALE",
            "description": "AUTHENTICTD SEC",
            "record_type": "detail",
            "interchange_direction": "CR",
        }
        self.env["bsr.vss.130.line"].create(
            [
                dict(common, region_country="A.P. - CANADA", fee_category="AUTHENTICTD SEC", interchange_amount=100.0),
                dict(common, region_country="A.P. - E.U.", fee_category="COMCL-BUS", interchange_amount=200.0),
            ]
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "The City",
                "product_type": "Reimbursement Fees",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "transaction_type": "ORIGINAL SALE",
                "business_mode": "ACQUIRER",
                "transaction_category": "PURCHASE",
                "description_filter": "AUTHENTICTD SEC",
            }
        )
        Option = self.env["bsr.accounting.filter.option"]
        region_options = Option.search(
            [("source_import_id", "=", self.vss_import.id), ("field_name", "=", "region_country")]
        )
        fee_options = Option.search(
            [("source_import_id", "=", self.vss_import.id), ("field_name", "=", "fee_category")]
        )
        self.assertIn("A.P. - CANADA", region_options.mapped("name"))
        self.assertIn("(Blank)", region_options.mapped("name"))
        self.assertIn("COMCL-BUS", fee_options.mapped("name"))
        self.assertIn("(Blank)", fee_options.mapped("name"))
        configuration.region_country_option_ids = region_options.filtered(
            lambda option: option.name == "A.P. - CANADA"
        )
        configuration.fee_category_option_ids = fee_options.filtered(
            lambda option: option.name == "AUTHENTICTD SEC"
        )
        configuration.action_match_source()
        self.assertEqual(len(configuration.line_ids), 1)
        self.assertEqual(configuration.line_ids.matched_amount, 100.0)

    def test_report_date_may_be_before_processing_date(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110"),
                "production_date": "2018-03-02",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "City Bank",
                "product_type": "Date ordering",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "settlement_currency": "BDT",
            }
        )
        self.assertEqual(configuration.report_date.isoformat(), "2018-03-01")
        self.assertEqual(configuration.production_date.isoformat(), "2018-03-02")

        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": self.vss_import.id,
                "production_date": "2018-03-02",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "product_type": "Date ordering",
                "configuration_ids": [(6, 0, configuration.ids)],
            }
        )
        self.assertEqual(run.report_date.isoformat(), "2018-03-01")

    def test_voucher_run_uses_selected_file_dates_without_changing_template(self):
        source_values = {
            "report_code": "VSS-120",
            "reporting_for_display": "DATE TEMPLATE",
            "rollup_to_display": "DATE ROLLUP",
            "funds_xfer_entity_display": "DATE FUNDS",
            "settlement_currency": "BDT",
            "clearing_currency": "BDT",
            "business_mode": "ISSUER",
            "transaction_category": "PURCHASE",
            "transaction_type": "ORIGINAL SALE",
            "clearing_direction": "DB",
            "record_type": "detail",
        }
        self.env["bsr.vss.120.line"].create(
            dict(
                source_values,
                import_id=self.vss_import.id,
                processing_date="2018-03-01",
                report_date="2018-03-01",
                clearing_amount=111.0,
            )
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Date Test",
                "product_type": "Interchange Value",
                "reporting_for": "DATE TEMPLATE",
                "rollup_to": "DATE ROLLUP",
                "funds_xfer_entity": "DATE FUNDS",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "vss_120_source_level": "detail",
                "amount_vss_120_interchange_credit": False,
                "amount_vss_120_interchange_debit": False,
            }
        )
        configuration.action_match_source()
        configuration.line_ids.write({"gl_number": "DATE-GL", "gl_name": "Date GL"})
        configuration.action_validate()

        attachment = self.env["ir.attachment"].create(
            {"name": "daily-date-file.txt", "raw": b"daily source", "mimetype": "text/plain"}
        )
        daily_import = self.env["bsr.vss.import"].create(
            {
                "attachment_id": attachment.id,
                "company_id": self.env.company.id,
                "state": "completed",
                "file_name": "daily-date-file.txt",
            }
        )
        self.env["bsr.vss.120.line"].create(
            dict(
                source_values,
                import_id=daily_import.id,
                processing_date="2026-09-15",
                report_date="2026-09-14",
                clearing_amount=222.0,
            )
        )

        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": daily_import.id,
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2026-09-26",
                "product_type": "Daily Interchange",
                "configuration_ids": [(6, 0, configuration.ids)],
            }
        )
        run.action_execute()

        self.assertEqual(run.production_date.isoformat(), "2026-09-15")
        self.assertEqual(run.report_date.isoformat(), "2026-09-14")
        self.assertEqual(configuration.production_date.isoformat(), "2018-03-01")
        self.assertEqual(configuration.report_date.isoformat(), "2018-03-01")
        runtime = run.generated_configuration_ids
        self.assertEqual(runtime.production_date.isoformat(), "2026-09-15")
        self.assertEqual(runtime.report_date.isoformat(), "2026-09-14")
        self.assertEqual(runtime.line_ids.matched_amount, 222.0)

    def test_vss_120_comma_filters_levels_and_all_amount_columns(self):
        common = {
            "import_id": self.vss_import.id,
            "report_code": "VSS-120",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "MULTI FILTER ENTITY",
            "rollup_to_display": "MULTI FILTER ROLLUP",
            "funds_xfer_entity_display": "MULTI FILTER FUNDS",
            "settlement_currency": "BDT",
            "clearing_currency": "BDT",
            "business_mode": "ACQUIRER",
            "disposition": "",
        }
        self.env["bsr.vss.120.line"].create(
            [
                dict(common, transaction_category="PURCHASE", transaction_type="ORIGINAL SALE",
                     clearing_amount=100.0, clearing_direction="DB", interchange_credit=5.0,
                     record_type="detail"),
                dict(common, transaction_category="ORIGINAL CREDIT", transaction_type="ORIGINAL",
                     clearing_amount=50.0, clearing_direction="CR", interchange_debit=2.0,
                     record_type="detail"),
                dict(common, transaction_category="PURCHASE", transaction_type="TOTAL PURCHASE",
                     clearing_amount=150.0, clearing_direction="DB", interchange_credit=7.0,
                     record_type="total"),
            ]
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "City Bank",
                "product_type": "VSS-120 filters",
                "reporting_for": "MULTI FILTER ENTITY",
                "rollup_to": "MULTI FILTER ROLLUP",
                "funds_xfer_entity": "MULTI FILTER FUNDS",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "Acquirer",
                "transaction_category": "PURCHASE, ORIGINAL CREDIT",
                "transaction_type": "ORIGINAL, ORIGINAL SALE",
                "vss_120_source_level": "detail",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 4)
        self.assertEqual(set(configuration.line_ids.mapped("record_type")), {"detail"})
        self.assertEqual(set(configuration.line_ids.mapped("source_currency")), {"BDT"})
        self.assertEqual(
            {(line.transaction_type, line.amount_source, line.source_direction, line.matched_amount)
             for line in configuration.line_ids},
            {
                ("ORIGINAL SALE", "clearing_amount", "DB", 100.0),
                ("ORIGINAL SALE", "interchange_credit", "CR", 5.0),
                ("ORIGINAL", "clearing_amount", "CR", 50.0),
                ("ORIGINAL", "interchange_debit", "DB", 2.0),
            },
        )

        configuration.action_reset_draft()
        configuration.amount_vss_120_clearing = False
        configuration.action_match_source()
        self.assertEqual(
            {(line.amount_source, line.matched_amount) for line in configuration.line_ids},
            {("interchange_credit", 5.0), ("interchange_debit", 2.0)},
        )

        configuration.action_reset_draft()
        configuration.write(
            {
                "transaction_category": "PURCHASE",
                "transaction_type": "TOTAL PURCHASE",
                "vss_120_source_level": "total",
                "amount_vss_120_clearing": True,
            }
        )
        configuration.action_match_source()
        self.assertEqual(len(configuration.line_ids), 2)
        self.assertEqual(set(configuration.line_ids.mapped("record_type")), {"total"})
        self.assertEqual(
            {(line.amount_source, line.source_direction, line.matched_amount) for line in configuration.line_ids},
            {("clearing_amount", "DB", 150.0), ("interchange_credit", "CR", 7.0)},
        )

    def test_vss_120_same_side_currency_pair_stays_on_one_gl_rule(self):
        self.env["bsr.vss.120.line"].create(
            {
                "import_id": self.vss_import.id,
                "report_code": "VSS-120",
                "processing_date": "2018-03-01",
                "report_date": "2018-03-01",
                "reporting_for_display": "PAIRING TEST",
                "rollup_to_display": "PAIRING TEST",
                "funds_xfer_entity_display": "PAIRING TEST",
                "settlement_currency": "USD",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "TOTAL ORIGINAL SALE",
                "disposition": "",
                "clearing_amount": 572110.45,
                "clearing_direction": "DB",
                "interchange_debit": 4549.40,
                "record_type": "total",
            }
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "UAT",
                "product_type": "VSS-120 pairing",
                "reporting_for": "PAIRING TEST",
                "rollup_to": "PAIRING TEST",
                "funds_xfer_entity": "PAIRING TEST",
                "settlement_currency": "USD",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "TOTAL ORIGINAL SALE",
                "source_description": "",
                "vss_120_source_level": "total",
                "amount_vss_120_interchange_credit": False,
            }
        )

        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 1)
        rule = configuration.line_ids
        self.assertEqual(rule.amount_source, "clearing_amount")
        self.assertEqual(rule.source_currency, "BDT")
        self.assertEqual(rule.posting_indicator, "DR")
        self.assertEqual(rule.matched_amount, 572110.45)
        self.assertEqual(rule.amount_columns_json, {"interchange_debit": 4549.40})
        self.assertEqual(configuration.preview_debit_total, 572110.45)
        self.assertEqual(configuration.preview_credit_total, 0.0)

        rule.write({"gl_number": "PAIR-GL", "gl_name": "Paired clearing GL"})
        configuration.action_reset_draft()
        configuration.action_match_source()
        self.assertEqual(len(configuration.line_ids), 1)
        self.assertEqual(configuration.line_ids.gl_number, "PAIR-GL")
        self.assertEqual(
            configuration.line_ids.amount_columns_json,
            {"interchange_debit": 4549.40},
        )

    def test_vss_210_respects_explicit_record_type_and_defaults_to_detail(self):
        common = {
            "import_id": self.vss_import.id,
            "report_code": "VSS-210",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "VSS-210 LEVEL TEST",
            "rollup_to_display": "VSS-210 LEVEL TEST",
            "funds_xfer_entity_display": "VSS-210 LEVEL TEST",
            "settlement_currency": "USD",
            "clearing_currency": "USD",
            "business_mode": "ISSUER",
            "transaction_category": "",
            "transaction_type": "",
            "program_network": "",
            "optional_issuer_fee_direction": "DB",
        }
        self.env["bsr.vss.210.line"].create(
            [
                dict(
                    common,
                    description="VISA INTERNATIONAL",
                    optional_issuer_fee=10.0,
                    record_type="detail",
                ),
                dict(
                    common,
                    description="TOTAL CURRENCY CONVERSION FEES",
                    optional_issuer_fee=91.14,
                    record_type="total",
                ),
            ]
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-210"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "UAT",
                "product_type": "VSS-210 record type",
                "reporting_for": "VSS-210 LEVEL TEST",
                "rollup_to": "VSS-210 LEVEL TEST",
                "funds_xfer_entity": "VSS-210 LEVEL TEST",
                "settlement_currency": "USD",
                "clearing_currency": "USD",
                "business_mode": "ISSUER",
                "amount_vss_210_settlement_interchange": False,
                "amount_vss_210_settlement_conversion": False,
                "amount_vss_210_clearing_interchange": False,
                "amount_vss_210_clearing_conversion": False,
                "amount_vss_210_optional_issuer": True,
            }
        )

        # A blank/unconfigured Record Type keeps DETAIL as the safe default.
        configuration.action_match_source()
        self.assertEqual(configuration.line_ids.record_type, "detail")
        self.assertEqual(configuration.line_ids.matched_amount, 10.0)

        configuration.action_reset_draft()
        configuration.write({"record_type": "total"})
        configuration.action_match_source()
        self.assertEqual(configuration.line_ids.record_type, "total")
        self.assertEqual(configuration.line_ids.description_match, "TOTAL CURRENCY CONVERSION FEES")
        self.assertEqual(configuration.line_ids.matched_amount, 91.14)

        configuration.action_reset_draft()
        configuration.write({"record_type": "detail"})
        configuration.action_match_source()
        self.assertEqual(configuration.line_ids.record_type, "detail")
        self.assertEqual(configuration.line_ids.matched_amount, 10.0)

    def test_only_selected_reports_are_posted(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Platinum Credit Card",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
            }
        )
        configuration.action_match_source()
        self.assertEqual(configuration.report_id_filter, "VSS-110, VSS-120")
        self.assertEqual(configuration.source_report_codes, "VSS-110, VSS-120")
        self.assertEqual(configuration.preview_debit_total, 946810.59)
        self.assertEqual(configuration.preview_credit_total, 946810.59)
        self.assertAlmostEqual(configuration.preview_difference, 0.0, places=2)

    def test_vss_110_alone_uses_all_non_duplicate_accounting_totals(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Settlement Summary Report",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                # VSS-110 has no Clearing Currency column, so this is ignored.
                "clearing_currency": "BDT",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 4)
        self.assertEqual(configuration.preview_debit_total, 946810.59)
        self.assertEqual(configuration.preview_credit_total, 946810.59)
        self.assertAlmostEqual(configuration.preview_difference, 0.0, places=2)
        self.assertEqual(
            {
                (
                    line.source_particulars,
                    line.source_entity,
                    line.amount_source,
                    line.source_direction,
                    line.posting_indicator,
                    line.matched_amount,
                )
                for line in configuration.line_ids
            },
            {
                ("INTERCHANGE VALUE", "TOTAL INTERCHANGE VALUE", "debit_amount", "DB", "DR", 946810.59),
                ("INTERCHANGE VALUE", "TOTAL INTERCHANGE VALUE", "credit_amount", "CR", "CR", 399.00),
                ("REIMBURSEMENT FEES", "TOTAL REIMBURSEMENT FEES", "total_amount", "CR", "CR", 16520.97),
                ("TOTAL", "NET SETTLEMENT AMOUNT", "total_amount", "DB", "CR", 929890.62),
            },
        )

    def test_vss_130_filters_pull_all_nonzero_detail_amounts(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "City Bank",
                "product_type": "Reimbursement Fees Report",
                "reporting_for": "1000548730 406591 SMS ACQ",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "business_mode": "Acquirer",
                "transaction_category": "ORIGINAL CREDIT",
                "transaction_type": "ORIGINAL",
                "description_filter": "ORIGINAL CREDIT",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 2)
        self.assertEqual(
            {
                (rule.record_type, rule.amount_source, rule.source_direction, rule.posting_indicator, rule.matched_amount)
                for rule in configuration.line_ids
            },
            {
                ("detail", "interchange_amount", "DB", "DR", 755191.65),
                ("detail", "fee_debit", "DB", "DR", 416.0),
            },
        )
        self.assertEqual(configuration.preview_debit_total, 755607.65)
        self.assertEqual(configuration.preview_credit_total, 0.0)

    def test_filtered_multi_report_does_not_duplicate_vss_130_interchange(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120", "VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Filtered multi-report voucher",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "business_mode": "Issuer",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 3)
        self.assertFalse(
            configuration.line_ids.filtered(
                lambda line: line.report_code == "VSS-130" and line.amount_source == "interchange_amount"
            )
        )
        self.assertEqual(configuration.preview_debit_total, 946810.59)
        self.assertEqual(configuration.preview_credit_total, 946418.97)

    def test_unfiltered_vss_130_acquirer_uses_one_aggregate_level(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "City Bank",
                "product_type": "Reimbursement Fees Report",
                "reporting_for": "1000548730 406591 SMS ACQ",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 1)
        line = configuration.line_ids
        self.assertEqual(line.description_match, "TOTAL ACQUIRER REIMB FEES")
        self.assertEqual(line.amount_source, "fee_debit")
        self.assertEqual(line.matched_amount, 416.0)


    def test_only_net_settlement_reverses_vss_110_direction(self):
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Selected VSS-110 entity",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "particulars": "INTERCHANGE VALUE",
                "entity": "TOTAL ISSUER",
            }
        )
        configuration.action_match_source()
        self.assertEqual(len(configuration.line_ids), 2)
        self.assertEqual(
            {
                (line.amount_source, line.source_direction, line.posting_indicator, line.matched_amount)
                for line in configuration.line_ids
            },
            {
                ("debit_amount", "DB", "DR", 946810.59),
                ("credit_amount", "CR", "CR", 399.00),
            },
        )
        self.assertEqual(set(configuration.line_ids.mapped("source_particulars")), {"INTERCHANGE VALUE"})
        self.assertEqual(set(configuration.line_ids.mapped("source_entity")), {"TOTAL ISSUER"})

        debit_rule = configuration.line_ids.filtered(lambda line: line.source_direction == "DB")
        debit_rule.posting_indicator = "CR"
        self.assertEqual(debit_rule.source_direction, "DB")
        self.assertEqual(debit_rule.posting_indicator, "CR")
        configuration.action_reset_draft()
        configuration.action_match_source()
        debit_rule = configuration.line_ids.filtered(lambda line: line.source_direction == "DB")
        self.assertEqual(debit_rule.posting_indicator, "CR")

        self.env["bsr.vss.110.line"].create(
            {
                "import_id": self.vss_import.id,
                "report_code": "VSS-110",
                "processing_date": "2018-03-01",
                "report_date": "2018-03-01",
                "reporting_for_display": "1000512940 BIN 474867 NNSS",
                "rollup_to_display": "1000137367 THE CITY DOM",
                "funds_xfer_entity_display": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "particulars": "TOTAL",
                "entity": "NET SETTLEMENT AMOUNT",
                "total_amount": 25.00,
                "direction": "CR",
                "record_type": "net",
            }
        )
        configuration.action_reset_draft()
        configuration.particulars = "TOTAL"
        configuration.entity = "NET SETTLEMENT AMOUNT"
        configuration.action_match_source()
        self.assertEqual(len(configuration.line_ids), 2)
        self.assertEqual(
            {(line.source_direction, line.posting_indicator) for line in configuration.line_ids},
            {("DB", "CR"), ("CR", "DR")},
        )

    def test_vss_130_multi_category_uses_all_requested_net_rows(self):
        common = {
            "import_id": self.vss_import.id,
            "report_code": "VSS-130",
            "processing_date": "2018-03-01",
            "report_date": "2018-03-01",
            "reporting_for_display": "1000137367 THE CITY DOM",
            "rollup_to_display": "1000137367 THE CITY DOM",
            "funds_xfer_entity_display": "1000137367 THE CITY DOM",
            "settlement_currency": "BDT",
            "business_mode": "ACQUIRER",
            "record_type": "net",
        }
        self.env["bsr.vss.130.line"].create(
            [
                dict(common, transaction_category="PURCHASE", transaction_type="CHARGEBACK",
                     description="NET PURCHASE", fee_debit=552174.71),
                dict(common, transaction_category="ATM CASH", transaction_type="CHARGEBACK",
                     description="NET ATM CASH", fee_credit=26675.00),
                dict(common, transaction_category="ORIGINAL CREDIT", transaction_type="ORIGINAL",
                     description="NET ORIGINAL CREDIT", fee_debit=416.00),
            ]
        )
        configuration = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "The City",
                "product_type": "Platinum",
                "reporting_for": "1000137367 THE CITY DOM",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "business_mode": "Acquirer",
                "transaction_category": "PURCHASE, ATM CASH, ORIGINAL CREDIT",
                "transaction_type": "CHARGEBACK",
                # This is now a disposition-only filter and VSS-130 has no
                # disposition column, so it must not suppress descriptions.
                "source_description": "NET PURCHASE, NET ATM CASH, NET ORIGINAL CREDIT",
            }
        )
        configuration.action_match_source()

        self.assertEqual(len(configuration.line_ids), 3)
        self.assertEqual(set(configuration.line_ids.mapped("record_type")), {"net"})
        self.assertEqual(
            {(line.description_match, line.posting_indicator, line.matched_amount) for line in configuration.line_ids},
            {
                ("NET PURCHASE", "DR", 552174.71),
                ("NET ATM CASH", "CR", 26675.00),
                ("NET ORIGINAL CREDIT", "DR", 416.00),
            },
        )

    def test_multi_configuration_run_generates_one_sequential_table(self):
        template = self.env["bsr.accounting.configuration"].create(
            {
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-110", "VSS-120", "VSS-130"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards Operation",
                "product_type": "Template Product",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
            }
        )
        template.action_match_source()
        template.line_ids.write({"gl_number": "99999999", "gl_name": "VISA outstanding GL"})
        template.action_validate()
        second_template = template.copy(default={"name": "BSR/CONFIG/TEST/SECOND"})
        second_template.action_validate()
        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": self.vss_import.id,
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "product_type": "Merged Product",
                "configuration_ids": [(6, 0, (template | second_template).ids)],
            }
        )

        run.action_execute()

        self.assertEqual(run.state, "done")
        self.assertEqual(len(run.generated_configuration_ids), 2)
        self.assertEqual(len(run.report_ids), 2)
        self.assertEqual(
            set(run.generated_configuration_ids.mapped("template_id").ids),
            {template.id, second_template.id},
        )
        self.assertEqual(
            set(run.generated_configuration_ids.mapped("product_type")),
            {"1000512940 BIN 474867 NNSS"},
        )
        self.assertAlmostEqual(run.debit_total, 1893621.18, places=2)
        self.assertAlmostEqual(run.credit_total, 1893621.18, places=2)
        workbook = load_workbook(io.BytesIO(base64.b64decode(run.output_file)), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Merged Voucher Report", "Merged Source Details"])
        sheet = workbook.active
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("1000512940 BIN 474867 NNSS", values)
        # Multiple configurations for one Reporting For value still produce
        # one voucher; per-product sections are only needed for multi-product runs.
        self.assertEqual(values.count("Debit Amount"), 1)
        self.assertEqual(values.count("Credit Amount"), 1)
        header_row = next(
            row[0].row
            for row in sheet.iter_rows()
            if row[0].value == "GL Number" and row[4].value == "GL Number"
        )
        self.assertEqual(
            [sheet.cell(header_row, column).value for column in range(1, 9)],
            [
                "GL Number", "GL Name", "Debit Amount", "Debit Remarks",
                "GL Number", "GL Name", "Credit Amount", "Credit Remarks",
            ],
        )
        total_row = next(row[0].row for row in sheet.iter_rows() if row[1].value == "Total")
        exported_amounts = [
            sheet.cell(row, column).value
            for row in range(header_row + 1, total_row)
            for column in (3, 7)
            if sheet.cell(row, column).value is not None
        ]
        self.assertEqual(len(exported_amounts), 8)
        self.assertEqual(exported_amounts.count(946810.59), 2)
        self.assertAlmostEqual(sheet.cell(total_row, 3).value, 1893621.18, places=2)
        self.assertAlmostEqual(sheet.cell(total_row, 7).value, 1893621.18, places=2)
        self.assertEqual(sheet.cell(total_row + 1, 6).value, "Difference")
        self.assertAlmostEqual(sheet.cell(total_row + 1, 7).value, 0.0, places=2)
        self.assertEqual(sheet.cell(total_row + 1, 8).value, "BALANCED")

    def test_run_product_types_come_from_reporting_for_and_create_individual_sections(self):
        attachment = self.env["ir.attachment"].create(
            {"name": "multi-product.txt", "raw": b"multi product", "mimetype": "text/plain"}
        )
        source_import = self.env["bsr.vss.import"].create(
            {
                "attachment_id": attachment.id,
                "company_id": self.env.company.id,
                "state": "completed",
                "file_name": "multi-product.txt",
            }
        )
        common = {
            "import_id": source_import.id,
            "report_code": "VSS-120",
            "processing_date": "2026-10-12",
            "report_date": "2026-10-12",
            "rollup_to_display": "MULTI ROLLUP",
            "funds_xfer_entity_display": "MULTI FUNDS",
            "settlement_currency": "BDT",
            "clearing_currency": "BDT",
            "business_mode": "ISSUER",
            "disposition": "",
            "record_type": "detail",
        }
        self.env["bsr.vss.120.line"].create(
            [
                dict(common, reporting_for_display="PRODUCT A", transaction_category="PURCHASE",
                     transaction_type="ORIGINAL SALE", clearing_direction="DB", clearing_amount=100.0),
                dict(common, reporting_for_display="PRODUCT A", transaction_category="MERCHANDISE CREDIT",
                     transaction_type="ORIGINAL", clearing_direction="CR", clearing_amount=100.0),
                dict(common, reporting_for_display="PRODUCT B", transaction_category="PURCHASE",
                     transaction_type="ORIGINAL SALE", clearing_direction="DB", clearing_amount=250.0),
                dict(common, reporting_for_display="PRODUCT B", transaction_category="MERCHANDISE CREDIT",
                     transaction_type="ORIGINAL", clearing_direction="CR", clearing_amount=250.0),
            ]
        )

        templates = self.env["bsr.accounting.configuration"]
        for product in ("PRODUCT A", "PRODUCT B"):
            template = self.env["bsr.accounting.configuration"].create(
                {
                    "source_import_id": source_import.id,
                    "report_code_ids": self._selected_reports("VSS-120"),
                    "production_date": "2026-10-12",
                    "report_date": "2026-10-12",
                    "posting_date": "2026-10-12",
                    "business_owner": "Cards",
                    "product_type": "Legacy manual value",
                    "reporting_for": product,
                    "rollup_to": "MULTI ROLLUP",
                    "funds_xfer_entity": "MULTI FUNDS",
                    "settlement_currency": "BDT",
                    "clearing_currency": "BDT",
                    "business_mode": "ISSUER",
                    "transaction_category": "PURCHASE, MERCHANDISE CREDIT",
                    "transaction_type": "ORIGINAL SALE, ORIGINAL",
                    "vss_120_source_level": "detail",
                    "amount_vss_120_interchange_credit": False,
                    "amount_vss_120_interchange_debit": False,
                }
            )
            template.action_match_source()
            template.line_ids.write(
                {"gl_number": "GL-%s" % product[-1], "gl_name": "%s clearing" % product}
            )
            template.action_validate()
            templates |= template

        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": source_import.id,
                "production_date": "2026-10-12",
                "report_date": "2026-10-12",
                "posting_date": "2026-10-12",
                "configuration_ids": [(6, 0, templates.ids)],
            }
        )

        self.assertEqual(run.product_type_selection_mode, "all")
        self.assertEqual(set(run.product_type_option_ids.mapped("name")), {"PRODUCT A", "PRODUCT B"})
        self.assertEqual(run.product_type, "PRODUCT A | PRODUCT B")
        self.assertTrue(templates <= run.available_configuration_ids)
        run.action_execute()

        self.assertEqual(run.state, "done")
        self.assertEqual(set(run.generated_configuration_ids.mapped("product_type")), {"PRODUCT A", "PRODUCT B"})
        workbook = load_workbook(io.BytesIO(base64.b64decode(run.output_file)), data_only=True)
        sheet = workbook["Merged Voucher Report"]
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertEqual(values.count("Debit Amount"), 3)
        self.assertIn("Product Voucher — PRODUCT A", values)
        self.assertIn("Product Voucher — PRODUCT B", values)

        product_a = run.product_type_option_ids.filtered(lambda option: option.name == "PRODUCT A")
        template_a = templates.filtered(lambda template: template.reporting_for == "PRODUCT A")
        single_run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": source_import.id,
                "production_date": "2026-10-12",
                "report_date": "2026-10-12",
                "posting_date": "2026-10-12",
                "product_type_option_ids": [(6, 0, product_a.ids)],
                "configuration_ids": [(6, 0, template_a.ids)],
            }
        )
        self.assertEqual(single_run.product_type_selection_mode, "custom")
        self.assertEqual(single_run.product_type, "PRODUCT A")
        self.assertIn(template_a, single_run.available_configuration_ids)
        single_run.action_execute()
        self.assertEqual(set(single_run.report_ids.configuration_id.mapped("reporting_for")), {"PRODUCT A"})
        single_values = [
            cell.value
            for row in load_workbook(
                io.BytesIO(base64.b64decode(single_run.output_file)), data_only=True
            )["Merged Voucher Report"].iter_rows()
            for cell in row
        ]
        self.assertNotIn("PRODUCT B", single_values)
        self.assertEqual(single_values.count("Debit Amount"), 1)
        self.assertNotIn("Product Voucher — PRODUCT A", single_values)

        custom_without_products = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": source_import.id,
                "production_date": "2026-10-12",
                "report_date": "2026-10-12",
                "posting_date": "2026-10-12",
                "product_type_selection_mode": "custom",
                "configuration_ids": [(6, 0, template_a.ids)],
            }
        )
        self.assertFalse(custom_without_products.product_type_option_ids)
        with self.assertRaisesRegex(UserError, "Product Type is set to Custom"):
            custom_without_products.action_execute()

        # One saved product's configuration set can be reused as the template
        # set for several Product Types selected from the daily VSS file.
        reusable_run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": source_import.id,
                "production_date": "2026-10-12",
                "report_date": "2026-10-12",
                "posting_date": "2026-10-12",
                "product_type_option_ids": [
                    (6, 0, run.product_type_option_ids.ids)
                ],
                "configuration_ids": [(6, 0, template_a.ids)],
            }
        )
        reusable_run.action_execute()
        self.assertEqual(
            set(reusable_run.report_ids.configuration_id.mapped("reporting_for")),
            {"PRODUCT A", "PRODUCT B"},
        )
        reusable_values = [
            cell.value
            for row in load_workbook(
                io.BytesIO(base64.b64decode(reusable_run.output_file)), data_only=True
            )["Merged Voucher Report"].iter_rows()
            for cell in row
        ]
        self.assertEqual(reusable_values.count("Debit Amount"), 3)
        self.assertIn("Product Voucher — PRODUCT A", reusable_values)
        self.assertIn("Product Voucher — PRODUCT B", reusable_values)

    def test_run_keeps_zero_activity_configuration_and_continues(self):
        common = {
            "report_code": "VSS-120",
            "reporting_for_display": "DAILY ENTITY",
            "rollup_to_display": "DAILY ROLLUP",
            "funds_xfer_entity_display": "DAILY FUNDS",
            "settlement_currency": "BDT",
            "clearing_currency": "BDT",
            "business_mode": "ISSUER",
            "disposition": "",
            "record_type": "detail",
        }

        def make_import(name):
            attachment = self.env["ir.attachment"].create(
                {"name": name, "raw": b"daily source", "mimetype": "text/plain"}
            )
            return self.env["bsr.vss.import"].create(
                {
                    "attachment_id": attachment.id,
                    "company_id": self.env.company.id,
                    "state": "completed",
                    "file_name": name,
                }
            )

        template_import = make_import("template-day.txt")
        daily_import = make_import("zero-activity-day.txt")
        self.env["bsr.vss.120.line"].create(
            [
                dict(common, import_id=template_import.id, processing_date="2026-09-19",
                     report_date="2026-09-19", transaction_category="PURCHASE",
                     transaction_type="ORIGINAL SALE", clearing_direction="DB", clearing_amount=100.0),
                dict(common, import_id=template_import.id, processing_date="2026-09-19",
                     report_date="2026-09-19", transaction_category="MERCHANDISE CREDIT",
                     transaction_type="ORIGINAL", clearing_direction="CR", clearing_amount=100.0),
                dict(common, import_id=template_import.id, processing_date="2026-09-19",
                     report_date="2026-09-19", transaction_category="ATM CASH",
                     transaction_type="ORIGINAL WITHDRAWAL", clearing_direction="DB", clearing_amount=50.0),
                dict(common, import_id=daily_import.id, processing_date="2026-09-20",
                     report_date="2026-09-20", transaction_category="PURCHASE",
                     transaction_type="ORIGINAL SALE", clearing_direction="DB", clearing_amount=125.0),
                dict(common, import_id=daily_import.id, processing_date="2026-09-20",
                     report_date="2026-09-20", transaction_category="MERCHANDISE CREDIT",
                     transaction_type="ORIGINAL", clearing_direction="CR", clearing_amount=125.0),
                dict(common, import_id=daily_import.id, processing_date="2026-09-20",
                     report_date="2026-09-20", transaction_category="ATM CASH",
                     transaction_type="ORIGINAL WITHDRAWAL", clearing_direction="DB", clearing_amount=0.0),
            ]
        )

        def make_template(name, category, transaction_type):
            configuration = self.env["bsr.accounting.configuration"].create(
                {
                    "name": name,
                    "source_import_id": template_import.id,
                    "report_code_ids": self._selected_reports("VSS-120"),
                    "production_date": "2026-09-19",
                    "report_date": "2026-09-19",
                    "posting_date": "2026-09-19",
                    "business_owner": "Cards",
                    "product_type": "Daily Product",
                    "reporting_for": "DAILY ENTITY",
                    "rollup_to": "DAILY ROLLUP",
                    "funds_xfer_entity": "DAILY FUNDS",
                    "settlement_currency": "BDT",
                    "clearing_currency": "BDT",
                    "business_mode": "ISSUER",
                    "transaction_category": category,
                    "transaction_type": transaction_type,
                    "vss_120_source_level": "detail",
                    "amount_vss_120_interchange_credit": False,
                    "amount_vss_120_interchange_debit": False,
                }
            )
            configuration.action_match_source()
            configuration.line_ids.write({"gl_number": "DAILY-GL", "gl_name": "Daily GL"})
            configuration.action_validate()
            return configuration

        active = make_template(
            "BSR/CONFIG/ACTIVE",
            "PURCHASE, MERCHANDISE CREDIT",
            "ORIGINAL SALE, ORIGINAL",
        )
        no_activity = make_template(
            "BSR/CONFIG/ZERO",
            "ATM CASH",
            "ORIGINAL WITHDRAWAL",
        )
        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": daily_import.id,
                "production_date": "2026-09-20",
                "report_date": "2026-09-20",
                "posting_date": "2026-09-20",
                "product_type": "Daily Product",
                "configuration_ids": [(6, 0, (active | no_activity).ids)],
            }
        )

        review_action = run.action_execute()
        self.assertEqual(review_action["res_model"], "bsr.accounting.run.review.wizard")
        review = self.env[review_action["res_model"]].browse(review_action["res_id"])
        self.assertEqual(review.warning_count, 1)
        self.assertEqual(review.blocking_count, 0)
        self.assertIn("zero-activity-day.txt", review.issue_summary_html)
        self.assertIn("BSR/CONFIG/ZERO", review.issue_summary_html)
        self.assertIn("Database ID", review.issue_summary_html)
        with self.assertRaisesRegex(UserError, "Acknowledge the listed issues"):
            review.action_continue()
        review.acknowledged = True
        review.action_continue()

        self.assertEqual(run.state, "done")
        self.assertTrue(run.issues_acknowledged)
        self.assertEqual(run.issues_acknowledged_by_id, self.env.user)
        self.assertTrue(run.issues_acknowledged_at)
        self.assertEqual(run.generated_report_count, 1)
        self.assertEqual(run.no_activity_count, 1)
        self.assertEqual(run.blocked_configuration_count, 0)
        self.assertEqual(len(run.generated_configuration_ids), 2)
        skipped = run.generated_configuration_ids.filtered(lambda config: config.state == "no_activity")
        self.assertEqual(skipped.template_id, no_activity)
        self.assertFalse(skipped.line_ids)
        self.assertIn("No Activity", skipped.validation_message)
        self.assertAlmostEqual(run.debit_total, 125.0, places=2)
        self.assertAlmostEqual(run.credit_total, 125.0, places=2)
        workbook = load_workbook(io.BytesIO(base64.b64decode(run.output_file)), data_only=True)
        details = workbook["Merged Source Details"]
        status_column = next(
            cell.column for cell in details[1] if cell.value == "Execution Status"
        )
        self.assertIn(
            "No Activity",
            [details.cell(row, status_column).value for row in range(2, details.max_row + 1)],
        )

    def test_run_reports_configuration_name_for_incompatible_daily_source(self):
        attachment = self.env["ir.attachment"].create(
            {"name": "incompatible-day.txt", "raw": b"daily source", "mimetype": "text/plain"}
        )
        daily_import = self.env["bsr.vss.import"].create(
            {
                "attachment_id": attachment.id,
                "company_id": self.env.company.id,
                "state": "completed",
                "file_name": "incompatible-day.txt",
            }
        )
        self.env["bsr.vss.120.line"].create(
            {
                "import_id": daily_import.id,
                "report_code": "VSS-120",
                "processing_date": "2026-09-20",
                "report_date": "2026-09-20",
                "reporting_for_display": "1000512940 BIN 474867 NNSS",
                "rollup_to_display": "1000137367 THE CITY DOM",
                "funds_xfer_entity_display": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "ATM CASH",
                "transaction_type": "ORIGINAL WITHDRAWAL",
                "clearing_direction": "DB",
                "clearing_amount": 10.0,
                "record_type": "detail",
            }
        )
        template = self.env["bsr.accounting.configuration"].create(
            {
                "name": "BSR/CONFIG/INCOMPATIBLE",
                "source_import_id": self.vss_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2018-03-01",
                "report_date": "2018-03-01",
                "posting_date": "2018-03-01",
                "business_owner": "Cards",
                "product_type": "Daily Product",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "PURCHASE",
                "transaction_type": "ORIGINAL SALE",
                "vss_120_source_level": "detail",
                "amount_vss_120_interchange_credit": False,
                "amount_vss_120_interchange_debit": False,
            }
        )
        template.action_match_source()
        template.line_ids.write({"gl_number": "DAILY-GL", "gl_name": "Daily GL"})
        template.action_validate()
        valid_template = self.env["bsr.accounting.configuration"].create(
            {
                "name": "BSR/CONFIG/COMPATIBLE",
                "source_import_id": daily_import.id,
                "report_code_ids": self._selected_reports("VSS-120"),
                "production_date": "2026-09-20",
                "report_date": "2026-09-20",
                "posting_date": "2026-09-20",
                "business_owner": "Cards",
                "product_type": "Daily Product",
                "reporting_for": "1000512940 BIN 474867 NNSS",
                "rollup_to": "1000137367 THE CITY DOM",
                "funds_xfer_entity": "1000137367 THE CITY DOM",
                "settlement_currency": "BDT",
                "clearing_currency": "BDT",
                "business_mode": "ISSUER",
                "transaction_category": "ATM CASH",
                "transaction_type": "ORIGINAL WITHDRAWAL",
                "vss_120_source_level": "detail",
                "amount_vss_120_interchange_credit": False,
                "amount_vss_120_interchange_debit": False,
            }
        )
        valid_template.action_match_source()
        valid_template.line_ids.write({"gl_number": "ATM-GL", "gl_name": "ATM GL"})
        valid_template.action_validate()
        run = self.env["bsr.accounting.run"].create(
            {
                "source_import_id": daily_import.id,
                "production_date": "2026-09-20",
                "report_date": "2026-09-20",
                "posting_date": "2026-09-20",
                "product_type": "Daily Product",
                "configuration_ids": [(6, 0, (template | valid_template).ids)],
            }
        )

        review_action = run.action_execute()
        self.assertEqual(review_action["res_model"], "bsr.accounting.run.review.wizard")
        review = self.env[review_action["res_model"]].browse(review_action["res_id"])
        self.assertEqual(review.warning_count, 0)
        self.assertEqual(review.blocking_count, 1)
        self.assertTrue(review.has_blocking_issues)
        self.assertIn("incompatible-day.txt", review.issue_summary_html)
        self.assertIn("BSR/CONFIG/INCOMPATIBLE", review.issue_summary_html)
        with self.assertRaisesRegex(UserError, "Acknowledge the listed issues"):
            review.action_continue()
        review.acknowledged = True
        review.action_continue()

        self.assertEqual(run.state, "exception")
        self.assertEqual(run.generated_report_count, 1)
        self.assertEqual(run.no_activity_count, 0)
        self.assertEqual(run.blocked_configuration_count, 1)
        self.assertTrue(run.issues_acknowledged)
        blocked = run.generated_configuration_ids.filtered(lambda config: config.state == "blocked")
        self.assertEqual(blocked.template_id, template)
        self.assertFalse(blocked.line_ids)
        self.assertIn("incompatible-day.txt", blocked.validation_message)
        self.assertIn("Database ID", blocked.validation_message)
        workbook = load_workbook(io.BytesIO(base64.b64decode(run.output_file)), data_only=True)
        details = workbook["Merged Source Details"]
        status_column = next(
            cell.column for cell in details[1] if cell.value == "Execution Status"
        )
        self.assertIn(
            "Skipped — Acknowledged Blocking Issue",
            [details.cell(row, status_column).value for row in range(2, details.max_row + 1)],
        )

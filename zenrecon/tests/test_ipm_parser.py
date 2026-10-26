# -*- coding: utf-8 -*-
import base64
import io
import json
import unittest

from openpyxl import load_workbook

from odoo.tests import tagged
from odoo.tests.common import TransactionCase

from ..services.ipm_parser import deblock, iter_rdw_records, mask_pan, parse_ipm, parse_iso8583


def _bitmap(*data_elements):
    value = bytearray(8)
    for de in data_elements:
        value[(de - 1) // 8] |= 0x80 >> ((de - 1) % 8)
    return bytes(value)


def _sample_physical_file():
    pan = b"5413331234567890"
    message = b"1240" + _bitmap(2, 3, 4, 24, 49) + b"16" + pan + b"000000" + b"000000012345" + b"200" + b"840"
    logical = len(message).to_bytes(4, "big") + message
    return logical + b" " * (1012 - len(logical)) + b"XX"


class TestIpmParser(unittest.TestCase):
    def test_reference_1240_record_and_pan_masking(self):
        document = parse_ipm(_sample_physical_file(), "sample.A14")
        self.assertEqual(document["file"]["record_count"], 1)
        self.assertEqual(document["file"]["error_count"], 0)
        record = document["records"][0]
        self.assertEqual(record["record_type"], "FIRST_PRESENTMENT")
        self.assertEqual(record["fields"]["2"]["value"], "541333******7890")
        self.assertNotIn("5413331234567890", json.dumps(document))
        self.assertEqual(record["fields"]["4"]["value"], "000000012345")

    def test_rdw_and_physical_block_rules(self):
        raw = _sample_physical_file()
        stream = deblock(raw)
        records = list(iter_rdw_records(stream))
        self.assertEqual(len(stream), 1012)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0][0], 1)
        self.assertNotIn(b"XX", records[0][2])

    def test_iso_and_short_pan_mask(self):
        message = b"1240" + _bitmap(3, 24) + b"000000" + b"200"
        parsed = parse_iso8583(message, "ascii")
        self.assertEqual((parsed["mti"], parsed["function_code"]), ("1240", "200"))
        self.assertEqual(mask_pan("1234567890"), "**********")


@tagged("post_install", "-at_install")
class TestIpmImport(TransactionCase):
    def test_file_is_stored_in_mastercard_tables_only(self):
        raw = _sample_physical_file()
        attachment = self.env["ir.attachment"].create({
            "name": "sample.A14", "raw": raw, "mimetype": "application/octet-stream",
        })
        batch = self.env["bsr.ipm.import"].create({"attachment_id": attachment.id})
        batch._process_one()
        self.assertEqual(batch.state, "completed")
        self.assertEqual(batch.record_count, 1)
        self.assertGreater(batch.field_count, 0)
        self.assertFalse(batch.error_count)
        record = batch.record_ids
        self.assertEqual(record.record_type, "FIRST_PRESENTMENT")
        self.assertEqual(record.pan_masked, "541333******7890")
        self.assertNotIn("5413331234567890", json.dumps(record.fields_json))
        self.assertFalse(self.env["bsr.vss.import"].search([("attachment_id", "=", attachment.id)]))

    def test_direct_upload_wizard_creates_attachment_and_import(self):
        wizard = self.env["bsr.ipm.import.wizard"].create({
            "upload_filename": "daily-mastercard.A14",
            "upload_file": base64.b64encode(_sample_physical_file()),
        })
        action = wizard.action_import_and_process()
        batch = self.env["bsr.ipm.import"].browse(action["res_id"])
        self.assertEqual(batch.state, "completed")
        self.assertEqual(batch.file_name, "daily-mastercard.A14")
        self.assertTrue(batch.attachment_id)
        self.assertEqual(batch.attachment_id.res_model, batch._name)
        self.assertEqual(batch.attachment_id.res_id, batch.id)

    def test_complete_xlsx_contains_all_data_elements(self):
        attachment = self.env["ir.attachment"].create({
            "name": "complete-export.A14", "raw": _sample_physical_file(),
            "mimetype": "application/octet-stream",
        })
        batch = self.env["bsr.ipm.import"].create({"attachment_id": attachment.id})
        batch._process_one()
        action = batch.action_download_complete_xlsx()
        wizard = self.env["bsr.ipm.export.wizard"].browse(action["res_id"])
        workbook = load_workbook(io.BytesIO(base64.b64decode(wizard.file_data)), data_only=True)
        self.assertEqual(workbook.sheetnames, ["File Metadata", "IPM Records", "Data Elements", "Parsing Errors"])
        records = workbook["IPM Records"]
        headers = [cell.value for cell in records[1]]
        self.assertIn("DE002", headers)
        self.assertIn("DE049", headers)
        self.assertEqual(records.cell(2, headers.index("DE002") + 1).value, "541333******7890")
        self.assertIsNone(records.cell(2, headers.index("Settlement Amount") + 1).value)
        self.assertEqual(workbook["Data Elements"].max_row, batch.field_count + 1)

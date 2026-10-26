# -*- coding: utf-8 -*-
import base64
from io import BytesIO

from openpyxl import load_workbook
from odoo.tests import tagged
from odoo.tests.common import TransactionCase
from odoo.exceptions import UserError


@tagged("post_install", "-at_install")
class TestNpsbSettlementVoucher(TransactionCase):
    def test_duplicate_pdf_cannot_be_processed_twice(self):
        content = base64.b64encode(b"same-npsb-pdf")
        original = self.env["npsb.settlement.report"].create({
            "pdf_file": content,
            "filename": "settlement.pdf",
            "state": "processed",
        })
        duplicate = self.env["npsb.settlement.report"].create({
            "pdf_file": content,
            "filename": "settlement-copy.pdf",
        })

        with self.assertRaisesRegex(UserError, original.display_name):
            duplicate.action_process_pdf()

    def test_dates_resolve_report_and_all_configurations_generate_voucher(self):
        report = self.env["npsb.settlement.report"].create({
            "pdf_file": base64.b64encode(b"test"),
            "filename": "daily-settlement.pdf",
            "state": "processed",
            "bank_name": "Dynamic Test Bank PLC",
            "transaction_date": "2026-10-08",
            "settlement_date": "2026-10-11",
        })
        self.env["npsb.settlement.summary.line"].create([
            {
                "report_id": report.id,
                "sequence": 1,
                "issuer_acquirer": "Issuing",
                "delivery_channel": "POS",
                "trans_type": "Purchase",
                "trans_amount": 1000.0,
            },
            {
                "report_id": report.id,
                "sequence": 2,
                "issuer_acquirer": "Acquiring/Beneficiary",
                "delivery_channel": "POS",
                "trans_type": "Purchase",
                "trans_amount": 1000.0,
            },
        ])
        template_report = self.env["npsb.settlement.report"].create({
            "pdf_file": base64.b64encode(b"template"),
            "filename": "configuration-source.pdf",
            "state": "processed",
            "transaction_date": "2026-09-08",
            "settlement_date": "2026-09-11",
        })
        self.env["npsb.settlement.summary.line"].create([
            {
                "report_id": template_report.id,
                "sequence": 1,
                "issuer_acquirer": "Issuing",
                "delivery_channel": "POS",
                "trans_type": "Purchase",
                "trans_amount": 900.0,
            },
            {
                "report_id": template_report.id,
                "sequence": 2,
                "issuer_acquirer": "Acquiring/Beneficiary",
                "delivery_channel": "POS",
                "trans_type": "Purchase",
                "trans_amount": 900.0,
            },
        ])
        configuration = self.env["npsb.settlement.configuration"].create({
            "source_report_id": template_report.id,
            "amount_trans_amount": True,
            "amount_fees_irf": False,
            "amount_platform_fee": False,
            "amount_net_amount": False,
        })
        options = self.env["npsb.settlement.filter.option"]
        option_domain = [("report_id", "=", template_report.id)]
        configuration.write({
            "transaction_date_option_ids": [(6, 0, options.search(
                option_domain + [("field_name", "=", "transaction_date")]
            ).ids)],
            "settlement_date_option_ids": [(6, 0, options.search(
                option_domain + [("field_name", "=", "settlement_date")]
            ).ids)],
            "delivery_channel_option_ids": [(6, 0, options.search(
                option_domain + [("field_name", "=", "delivery_channel"), ("value_key", "=", "pos")]
            ).ids)],
            "transaction_type_option_ids": [(6, 0, options.search(
                option_domain + [("field_name", "=", "transaction_type"), ("value_key", "=", "purchase")]
            ).ids)],
        })
        self.env["npsb.settlement.configuration.line"].create([
            {
                "configuration_id": configuration.id,
                "sequence": 10,
                "business_mode": "Issuing",
                "delivery_channel": "POS",
                "transaction_type": "Purchase",
                "amount_source": "trans_amount",
                "matched_amount": 1000.0,
                "include_in_voucher": True,
                "posting_indicator": "DR",
                "gl_number": "110100",
                "gl_name": "NPSB Settlement Receivable",
                "debit_remarks": "NPSB issuing settlement",
            },
            {
                "configuration_id": configuration.id,
                "sequence": 20,
                "business_mode": "Acquiring/Beneficiary",
                "delivery_channel": "POS",
                "transaction_type": "Purchase",
                "amount_source": "trans_amount",
                "matched_amount": 1000.0,
                "include_in_voucher": True,
                "posting_indicator": "CR",
                "gl_number": "210100",
                "gl_name": "NPSB Settlement Payable",
                "credit_remarks": "NPSB acquiring settlement",
            },
        ])

        # The configuration intentionally remains Draft and was authored from
        # another day's report with its source dates selected. Voucher dates
        # choose the daily report; reusable business filters still apply.
        self.assertEqual(configuration.state, "draft")
        self.assertNotEqual(configuration.source_report_id, report)
        self.assertTrue(configuration.transaction_date_option_ids)
        self.assertTrue(configuration.settlement_date_option_ids)
        run = self.env["npsb.settlement.voucher.run"].create({
            "transaction_date": "2026-10-08",
            "settlement_date": "2026-10-11",
        })
        run.action_generate()

        self.assertEqual(run.source_report_id, report)
        self.assertEqual(run.configuration_ids, configuration)
        self.assertEqual(run.state, "generated")
        self.assertEqual(len(run.line_ids), 2)
        self.assertEqual(run.debit_total, 1000.0)
        self.assertEqual(run.credit_total, 1000.0)
        self.assertTrue(run.is_balanced)
        output = base64.b64decode(run.output_file)
        self.assertTrue(output.startswith(b"PK"))
        workbook = load_workbook(BytesIO(output), data_only=True)
        self.assertEqual(workbook.sheetnames, ["NPSB Voucher", "Audit Details"])
        voucher = workbook["NPSB Voucher"]
        self.assertEqual(voucher["A1"].value, "BANK RECONCILIATION VOUCHER")
        self.assertEqual(voucher["A2"].value, "Dynamic Test Bank PLC")
        self.assertEqual(voucher["A9"].value, "↓  DEBIT ENTRIES")
        self.assertEqual(voucher["G9"].value, "↑  CREDIT ENTRIES")

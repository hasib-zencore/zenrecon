# -*- coding: utf-8 -*-
import base64
import html
import io

from odoo import api, fields, models, _
from odoo.exceptions import UserError

from .accounting_voucher import (
    AMOUNT_SOURCE_LABELS,
    AMOUNT_SOURCE_ORDER,
    _currency_balance_adjustments,
    _group_source_details,
    _mapping_amount_currency,
    _mapping_currency_values,
    _paired_amount_field,
    _voucher_row_height,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None


REPORT_ORDER = {f"VSS-{code}": index for index, code in enumerate(("110", "120", "130", "140", "210"))}
SOURCE_LINE_MODELS = tuple(
    f"bsr.vss.{code}.line" for code in ("110", "120", "130", "140", "210")
)


class BsrAccountingProductOption(models.Model):
    _name = "bsr.accounting.product.option"
    _description = "Processed VSS Reporting For / Product Type"
    _order = "name, id"

    source_import_id = fields.Many2one(
        "bsr.vss.import", required=True, ondelete="cascade", index=True
    )
    company_id = fields.Many2one(
        related="source_import_id.company_id", store=True, index=True
    )
    name = fields.Char(string="Product Type", required=True, index=True)
    value_key = fields.Char(required=True, index=True)

    _source_value_unique = models.Constraint(
        "UNIQUE(source_import_id, value_key)",
        "Each Reporting For value must be unique within its processed VSS file.",
    )

    @staticmethod
    def _value_key(value):
        return " ".join(str(value or "").strip().casefold().split())

    @api.model
    def ensure_for_import(self, source_import):
        """Materialize the unique Reporting For values parsed from one file."""
        if not source_import:
            return self.browse()
        values = {}
        for model_name in SOURCE_LINE_MODELS:
            lines = self.env[model_name].sudo().search(
                [("import_id", "=", source_import.id)]
            )
            for value in lines.mapped("reporting_for_display"):
                cleaned = " ".join(str(value or "").strip().split())
                key = self._value_key(cleaned)
                if key:
                    values.setdefault(key, cleaned)
        existing = self.sudo().search([("source_import_id", "=", source_import.id)])
        known = set(existing.mapped("value_key"))
        create_values = [
            {
                "source_import_id": source_import.id,
                "name": name,
                "value_key": key,
            }
            for key, name in values.items()
            if key not in known
        ]
        if create_values:
            self.sudo().create(create_values)
        return self.search(
            [
                ("source_import_id", "=", source_import.id),
                ("value_key", "in", list(values)),
            ]
        )


class BsrAccountingRun(models.Model):
    _name = "bsr.accounting.run"
    _description = "Multi-Configuration Voucher Run"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "create_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(default=lambda self: _("New"), readonly=True, copy=False, tracking=True)
    company_id = fields.Many2one(
        "res.company", required=True, default=lambda self: self.env.company, index=True, tracking=True
    )
    source_import_id = fields.Many2one(
        "bsr.vss.import",
        string="Processed VSS File",
        index=True,
        ondelete="restrict",
        check_company=True,
        domain="[('company_id', '=', company_id), ('state', 'in', ('completed', 'partial'))]",
        tracking=True,
    )
    production_date = fields.Date(string="PROC Date", required=True, tracking=True)
    report_date = fields.Date(required=True, tracking=True)
    posting_date = fields.Date(required=True, default=fields.Date.context_today, tracking=True)
    business_owner_filter_id = fields.Many2one(
        "bsr.visa.dimension.option",
        string="Business Owner",
        domain="[('company_id', '=', company_id), ('dimension', '=', 'business_owner')]",
        tracking=True,
    )
    product_type_filter_id = fields.Many2one(
        "bsr.visa.dimension.option",
        string="Product Type",
        domain="[('company_id', '=', company_id), ('dimension', '=', 'product_type')]",
        tracking=True,
    )
    card_type_filter_id = fields.Many2one(
        "bsr.visa.dimension.option",
        string="Card Type",
        domain="[('company_id', '=', company_id), ('dimension', '=', 'card_type')]",
        tracking=True,
    )
    settlement_currency_scope = fields.Selection(
        [("all", "All"), ("usd", "USD"), ("bdt", "BDT")],
        string="Settlement Currency",
        required=True,
        default="all",
        tracking=True,
    )
    visa_setup_ids = fields.Many2many(
        "bsr.visa.setup",
        "bsr_accounting_run_visa_setup_rel",
        "run_id",
        "setup_id",
        string="Resolved VISA Setups",
        readonly=True,
        copy=False,
    )
    product_type_selection_mode = fields.Selection(
        [("all", "All"), ("custom", "Custom")],
        string="Product Type",
        required=True,
        default="all",
        tracking=True,
        help=(
            "All includes every unique Reporting For value in the processed VSS file. "
            "Custom allows one or more values to be selected."
        ),
    )
    product_type = fields.Char(
        string="Included Product Types",
        readonly=True,
        tracking=True,
        help="Audit display of the selected Reporting For values from the processed VSS file.",
    )
    product_type_option_ids = fields.Many2many(
        "bsr.accounting.product.option",
        "bsr_accounting_run_product_option_rel",
        "run_id",
        "option_id",
        string="Selected Product Types",
        tracking=True,
        help="Select one or more unique Reporting For values found in the processed VSS file.",
    )
    available_configuration_ids = fields.Many2many(
        "bsr.accounting.configuration",
        compute="_compute_available_configuration_ids",
        string="Available Configuration Files",
    )
    usd_buy_rate = fields.Float(string="USD Buy Rate", digits=(20, 6), tracking=True)
    usd_sale_rate = fields.Float(string="USD Sale Rate", digits=(20, 6), tracking=True)
    usd_mid_rate = fields.Float(string="USD Mid Rate", digits=(20, 6), tracking=True)
    currency_conversion_gl_number = fields.Char(
        string="Currency Conversion GL",
        tracking=True,
        help="Balancing GL used for currency conversion gain or loss.",
    )
    configuration_ids = fields.Many2many(
        "bsr.accounting.configuration",
        "bsr_accounting_run_configuration_rel",
        "run_id",
        "configuration_id",
        string="Configuration Files",
        check_company=True,
        domain="[('company_id', '=', company_id), ('state', '=', 'ready'), ('batch_run_id', '=', False)]",
        help="Select one or more completed configuration files. Each is matched independently and printed in source-report order.",
        tracking=True,
    )
    requested_report_ids = fields.Char(compute="_compute_requested_reports", string="Report IDs")
    generated_configuration_ids = fields.One2many(
        "bsr.accounting.configuration", "batch_run_id", string="Executed Configurations", readonly=True
    )
    report_ids = fields.One2many("bsr.accounting.report", "batch_run_id", string="Generated Reports", readonly=True)
    generated_report_count = fields.Integer(compute="_compute_execution_counts")
    no_activity_count = fields.Integer(compute="_compute_execution_counts")
    blocked_configuration_count = fields.Integer(compute="_compute_execution_counts")
    issues_acknowledged = fields.Boolean(readonly=True, copy=False, tracking=True)
    issues_acknowledged_at = fields.Datetime(readonly=True, copy=False)
    issues_acknowledged_by_id = fields.Many2one(
        "res.users", string="Issues Acknowledged By", readonly=True, copy=False
    )
    debit_total = fields.Float(compute="_compute_totals", digits=(20, 2))
    credit_total = fields.Float(compute="_compute_totals", digits=(20, 2))
    difference = fields.Float(compute="_compute_totals", digits=(20, 2))
    is_balanced = fields.Boolean(compute="_compute_totals")
    state = fields.Selection(
        [("draft", "Draft"), ("done", "Generated"), ("exception", "Exception")],
        required=True,
        default="draft",
        index=True,
        tracking=True,
    )
    output_file = fields.Binary(attachment=True, readonly=True, copy=False)
    output_filename = fields.Char(readonly=True, copy=False)
    generated_at = fields.Datetime(readonly=True, copy=False)
    generated_by_id = fields.Many2one("res.users", readonly=True, copy=False, string="Generated By")

    @api.model
    def default_get(self, field_list):
        values = super().default_get(field_list)
        company = self.env["res.company"].browse(
            values.get("company_id") or self.env.company.id
        )
        Option = self.env["bsr.visa.dimension.option"]
        defaults = (
            ("business_owner_filter_id", "business_owner"),
            ("product_type_filter_id", "product_type"),
            ("card_type_filter_id", "card_type"),
        )
        for field_name, dimension in defaults:
            if field_name in field_list and not values.get(field_name):
                values[field_name] = Option.all_option(company, dimension).id
        return values

    @api.model_create_multi
    def create(self, vals_list):
        for values in vals_list:
            if not values.get("name") or values.get("name") == _("New"):
                values["name"] = self.env["ir.sequence"].next_by_code("bsr.accounting.run") or _("New")
        explicit_scope = ["product_type_selection_mode" in values for values in vals_list]
        explicit_products = ["product_type_option_ids" in values for values in vals_list]
        records = super().create(vals_list)
        for run, has_scope, has_products in zip(records, explicit_scope, explicit_products):
            if not run.source_import_id:
                continue
            available = run._available_product_type_options()
            if not has_scope:
                selected = run.product_type_option_ids
                if not has_products and run.configuration_ids:
                    configuration_keys = {
                        BsrAccountingProductOption._value_key(value)
                        for value in run.configuration_ids.mapped("reporting_for")
                        if value
                    }
                    selected = available.filtered(
                        lambda option: option.value_key in configuration_keys
                    )
                # Preserve legacy/API callers that supplied or implied a
                # subset before the selector existed. Supplying every option
                # still means All. The form itself explicitly sends All.
                if selected and set(selected.ids) != set(available.ids):
                    run.with_context(skip_product_type_scope_sync=True).write(
                        {
                            "product_type_selection_mode": "custom",
                            "product_type_option_ids": [(6, 0, selected.ids)],
                        }
                    )
            run._synchronize_product_type_scope()
        return records

    def write(self, values):
        result = super().write(values)
        if self.env.context.get("skip_product_type_scope_sync"):
            return result
        scope_changed = "product_type_selection_mode" in values
        source_changed = "source_import_id" in values
        if scope_changed or source_changed:
            clear_custom = (
                "product_type_option_ids" not in values
                and any(run.product_type_selection_mode == "custom" for run in self)
            )
            self._synchronize_product_type_scope(clear_custom=clear_custom)
        elif "product_type_option_ids" in values:
            self._synchronize_product_type_scope()
        return result

    def _available_product_type_options(self):
        self.ensure_one()
        if not self.source_import_id:
            return self.env["bsr.accounting.product.option"]
        return self.env["bsr.accounting.product.option"].ensure_for_import(
            self.source_import_id
        )

    def _synchronize_product_type_scope(self, clear_custom=False):
        """Keep the existing execution relation aligned with All/Custom."""
        for run in self:
            available = run._available_product_type_options()
            if run.product_type_selection_mode == "all":
                selected = available
            elif clear_custom:
                selected = self.env["bsr.accounting.product.option"]
            else:
                selected = run.product_type_option_ids & available
            if set(run.product_type_option_ids.ids) != set(selected.ids):
                run.with_context(skip_product_type_scope_sync=True).write(
                    {"product_type_option_ids": [(6, 0, selected.ids)]}
                )
            run._sync_product_type_display()

    def _sync_product_type_display(self):
        for run in self:
            display = " | ".join(run.product_type_option_ids.sorted("name").mapped("name"))
            if run.product_type != display:
                # Direct assignment also works safely on unsaved onchange records.
                run.product_type = display or False

    @api.depends("company_id", "source_import_id")
    def _compute_available_configuration_ids(self):
        Configuration = self.env["bsr.accounting.configuration"]
        for run in self:
            if not run.source_import_id:
                run.available_configuration_ids = False
                continue
            # Configuration files are reusable accounting templates. Their
            # saved Reporting For value identifies the original product, but
            # a selected template set may be applied to other Product Types
            # chosen from the daily processed file.
            run.available_configuration_ids = Configuration.search(
                [
                    ("company_id", "=", run.company_id.id),
                    ("state", "=", "ready"),
                    ("batch_run_id", "=", False),
                ]
            )

    def _conversion_rate_for(self, conversion_type):
        self.ensure_one()
        return {"buy": self.usd_buy_rate, "sale": self.usd_sale_rate, "mid": self.usd_mid_rate}.get(conversion_type or "mid", 0.0)

    @api.depends("configuration_ids.report_id_filter")
    def _compute_requested_reports(self):
        standard = ("VSS-110", "VSS-120", "VSS-130", "VSS-140", "VSS-210")
        for run in self:
            selected = set(run.configuration_ids.report_code_ids.mapped("code"))
            run.requested_report_ids = ", ".join(code for code in standard if code in selected)

    @api.depends("generated_configuration_ids.state", "report_ids")
    def _compute_execution_counts(self):
        for run in self:
            run.generated_report_count = len(run.report_ids)
            run.no_activity_count = len(
                run.generated_configuration_ids.filtered(lambda config: config.state == "no_activity")
            )
            run.blocked_configuration_count = len(
                run.generated_configuration_ids.filtered(lambda config: config.state == "blocked")
            )

    @api.depends(
        "report_ids.debit_total",
        "report_ids.credit_total",
        "generated_configuration_ids.state",
    )
    def _compute_totals(self):
        for run in self:
            if not run.report_ids:
                run.debit_total = 0.0
                run.credit_total = 0.0
                run.difference = 0.0
                run.is_balanced = bool(run.generated_configuration_ids) and all(
                    config.state == "no_activity" for config in run.generated_configuration_ids
                )
                continue
            ordered_reports = run.report_ids.sorted(
                lambda report: run._template_sort_key(report.configuration_id.template_id)
            )
            data = run._voucher_section_data(run._merged_mapping_groups(ordered_reports))
            run.debit_total = data["debit_bdt"]
            run.credit_total = data["credit_bdt"]
            run.difference = data["difference"]
            run.is_balanced = data["difference"] < 0.005

    @api.onchange("source_import_id")
    def _onchange_source_import_id(self):
        for run in self:
            run.configuration_ids = False
            if run.source_import_id and run.product_type_selection_mode == "all":
                run.product_type_option_ids = run._available_product_type_options()
            else:
                run.product_type_option_ids = False
            run._sync_product_type_display()
            dates = run._source_import_dates()
            if dates[0]:
                run.production_date = dates[0]
            if dates[1]:
                run.report_date = dates[1]

    @api.onchange("company_id")
    def _onchange_company_id_visa_filters(self):
        Option = self.env["bsr.visa.dimension.option"]
        for run in self:
            run.business_owner_filter_id = Option.all_option(
                run.company_id, "business_owner"
            )
            run.product_type_filter_id = Option.all_option(
                run.company_id, "product_type"
            )
            run.card_type_filter_id = Option.all_option(run.company_id, "card_type")

    @api.onchange("product_type_selection_mode")
    def _onchange_product_type_selection_mode(self):
        for run in self:
            if run.product_type_selection_mode == "all":
                run.product_type_option_ids = run._available_product_type_options()
            else:
                run.product_type_option_ids = False
            run._sync_product_type_display()

    @api.onchange("product_type_option_ids")
    def _onchange_product_type_option_ids(self):
        for run in self:
            run._sync_product_type_display()
            run.configuration_ids &= run.available_configuration_ids

    def _source_import_dates(self):
        """Return the unique PROC and report dates carried by the selected file."""
        self.ensure_one()
        if not self.source_import_id:
            return False, False
        production_dates, report_dates = self._source_import_date_values()
        return (
            next(iter(production_dates)) if len(production_dates) == 1 else False,
            next(iter(report_dates)) if len(report_dates) == 1 else False,
        )

    def _configuration_matches_currency_scope(self, configuration):
        """A group's ALL scope never changes the configuration's own currency."""
        self.ensure_one()
        if self.settlement_currency_scope == "all":
            return True
        configuration_currency = (configuration.settlement_currency or "").strip().upper()
        return configuration_currency == self.settlement_currency_scope.upper()

    def _validate_selected_configuration_currencies(self):
        """Block explicit/API selections that bypass automatic group filtering."""
        self.ensure_one()
        if self.settlement_currency_scope == "all":
            return
        incompatible = self.configuration_ids.filtered(
            lambda configuration: not self._configuration_matches_currency_scope(configuration)
        )
        if not incompatible:
            return
        details = ", ".join(
            "%s (%s)" % (
                configuration.display_name,
                (configuration.settlement_currency or _("Blank")).strip().upper(),
            )
            for configuration in incompatible[:10]
        )
        if len(incompatible) > 10:
            details += _(", and %s more") % (len(incompatible) - 10)
        raise UserError(
            _(
                "Settlement Currency mismatch. This voucher run is %(run_currency)s, but the "
                "following selected configuration(s) are not %(run_currency)s: %(configurations)s. "
                "A group marked ALL is available to both run types, but it does not override each "
                "configuration's own settlement currency."
            )
            % {
                "run_currency": self.settlement_currency_scope.upper(),
                "configurations": details,
            }
        )

    def _prepare_visa_run_scope(self):
        """Resolve the simple run inputs to the existing execution inputs.

        Explicit legacy/API runs remain supported: once a source and
        configurations are supplied, this orchestration layer does nothing.
        """
        self.ensure_one()
        if self.source_import_id and self.configuration_ids:
            return
        if not self.production_date or not self.report_date:
            raise UserError(_("Select both Production Date and Report Date."))

        candidate_import_ids = set()
        line_domain = [
            ("import_id.company_id", "=", self.company_id.id),
            ("import_id.state", "in", ("completed", "partial")),
            ("processing_date", "=", self.production_date),
            ("report_date", "=", self.report_date),
        ]
        for model_name in SOURCE_LINE_MODELS:
            candidate_import_ids.update(
                self.env[model_name].sudo().search(line_domain).mapped("import_id").ids
            )
        source_import = self.env["bsr.vss.import"].search(
            [("id", "in", list(candidate_import_ids))],
            order="processing_completed_at desc, id desc",
            limit=1,
        )
        if not source_import:
            raise UserError(
                _(
                    "No successfully processed VSS file contains Production Date %(production)s "
                    "and Report Date %(report)s."
                )
                % {
                    "production": self.production_date,
                    "report": self.report_date,
                }
            )

        # The resolved daily VSS file is authoritative for VISA Setup source
        # identities. Friendly Business Owner/Product/Card naming remains on
        # the permanent setup records and is reused whenever an identity returns.
        self.env["bsr.visa.source.option"].sync_from_import(source_import)

        setup_domain = [("company_id", "=", self.company_id.id), ("active", "=", True)]
        dimensions = (
            (self.business_owner_filter_id, "business_owner", _("Business Owner")),
            (self.product_type_filter_id, "product_type", _("Product Type")),
            (self.card_type_filter_id, "card_type", _("Card Type")),
        )
        for option, field_name, label in dimensions:
            if not option:
                raise UserError(_("Select %s or All.") % label)
            if not option.is_all:
                setup_domain.append((field_name, "=ilike", option.name))
        setups = self.env["bsr.visa.setup"].search(setup_domain)
        if not setups:
            raise UserError(
                _("No VISA Setup mapping matches the selected Business Owner, Product Type, and Card Type.")
            )

        mapped_keys = set(setups.reporting_for_option_ids.mapped("value_key"))
        available_products = self.env["bsr.accounting.product.option"].ensure_for_import(
            source_import
        )
        products = available_products.filtered(lambda option: option.value_key in mapped_keys)
        if not products:
            raise UserError(
                _(
                    "The selected VISA Setup Reporting For values are not present in the processed "
                    "VSS file for these dates."
                )
            )

        # Exclude setup rows which only reference identities from older files.
        current_product_keys = set(products.mapped("value_key"))
        setups = setups.filtered(
            lambda setup: bool(
                current_product_keys
                & set(setup.reporting_for_option_ids.filtered("active").mapped("value_key"))
            )
        )

        currency_domain = {
            "usd": ("usd", "all"),
            "bdt": ("bdt", "all"),
            "all": ("usd", "bdt", "all"),
        }[self.settlement_currency_scope]
        groups = self.env["bsr.accounting.configuration.group"].search(
            [
                ("company_id", "=", self.company_id.id),
                ("settlement_currency", "in", currency_domain),
            ]
        )
        candidate_configurations = self.env["bsr.accounting.configuration"].search(
            [
                ("company_id", "=", self.company_id.id),
                ("state", "=", "ready"),
                ("batch_run_id", "=", False),
                ("configuration_group_id", "in", groups.ids),
            ]
        )
        configurations = candidate_configurations.filtered(
            lambda configuration: self._configuration_matches_currency_scope(configuration)
        )
        if not configurations:
            if candidate_configurations and self.settlement_currency_scope != "all":
                found_currencies = ", ".join(
                    sorted({
                        (configuration.settlement_currency or _("Blank")).strip().upper()
                        for configuration in candidate_configurations
                    })
                )
                raise UserError(
                    _(
                        "Settlement Currency mismatch. The selected %(run_currency)s run found "
                        "eligible VISA group(s), but their ready configurations use %(found)s. "
                        "Assign at least one %(run_currency)s configuration to a compatible group."
                    )
                    % {
                        "run_currency": self.settlement_currency_scope.upper(),
                        "found": found_currencies,
                    }
                )
            raise UserError(
                _("No ready VISA Configuration group is assigned to the selected settlement currency.")
            )

        self.with_context(skip_product_type_scope_sync=True).write(
            {
                "source_import_id": source_import.id,
                "configuration_ids": [(6, 0, configurations.ids)],
                "product_type_selection_mode": "custom",
                "product_type_option_ids": [(6, 0, products.ids)],
                "visa_setup_ids": [(6, 0, setups.ids)],
            }
        )
        self._sync_product_type_display()

    def _source_import_date_values(self):
        """Return all date choices available in the selected processed file."""
        self.ensure_one()
        production_dates = set()
        report_dates = set()
        for model_name in SOURCE_LINE_MODELS:
            lines = self.env[model_name].search([("import_id", "=", self.source_import_id.id)])
            production_dates.update(filter(None, lines.mapped("processing_date")))
            report_dates.update(filter(None, lines.mapped("report_date")))
        return production_dates, report_dates

    def _apply_source_import_dates(self):
        """Make the selected file's dates authoritative for this voucher run."""
        self.ensure_one()
        production_date, report_date = self._source_import_dates()
        available_production, available_reports = self._source_import_date_values()
        production_date = production_date or self.production_date
        report_date = report_date or self.report_date
        if (
            not production_date
            or not report_date
            or production_date not in available_production
            or report_date not in available_reports
        ):
            raise UserError(
                _(
                    "Select a PROC Date and Report Date available in the processed VSS file."
                )
            )
        self.write({"production_date": production_date, "report_date": report_date})
        return production_date, report_date

    @staticmethod
    def _template_sort_key(configuration):
        codes = configuration.report_code_ids.mapped("code")
        first = min((REPORT_ORDER.get(code, 99) for code in codes), default=99)
        return first, configuration.id

    def _merged_mapping_groups(self, ordered_reports):
        """Use only rules re-matched against this run's selected processed file."""
        groups = []
        for report in ordered_reports:
            mappings = report.configuration_id.line_ids.filtered("include_in_voucher").sorted(
                lambda line: (line.sequence, line.id)
            )
            groups.append((report, mappings))
        return groups

    def _validate_currency_setup(self, configuration):
        """Validate only conversion types and rates actually used by USD rules."""
        self.ensure_one()
        usd_rules = configuration.line_ids.filtered(
            lambda line: line.include_in_voucher
            and (line.source_currency or "").strip().upper() == "USD"
            and not _paired_amount_field(line)
        )
        if not usd_rules:
            return
        if not self.currency_conversion_gl_number:
            raise UserError(_("Enter the Currency Conversion GL before generating the voucher."))
        missing_type = usd_rules.filtered(lambda line: not line.conversion_type)
        if missing_type:
            raise UserError(
                _("Select a Conversion Type for every USD accounting rule in configuration %s.")
                % (configuration.template_id.name or configuration.name)
            )
        missing_rates = sorted(
            conversion_type
            for conversion_type in set(usd_rules.mapped("conversion_type"))
            if self._conversion_rate_for(conversion_type) <= 0
        )
        if missing_rates:
            labels = dict(usd_rules._fields["conversion_type"].selection)
            raise UserError(
                _("Enter a positive rate for: %s")
                % ", ".join(labels.get(rate, rate) for rate in missing_rates)
            )

    def _execution_plan(self):
        """Pair selected configuration templates with selected Product Types.

        A configuration selection containing one original Reporting For value
        is treated as one reusable template set and is executed for every
        selected Product Type. If the selection already contains several
        product-specific sets, each set is executed only for its matching
        Product Type so mappings are never duplicated across products.
        """
        self.ensure_one()
        options = self.product_type_option_ids.sorted("name")
        templates = self.configuration_ids.sorted(self._template_sort_key)
        template_keys = {
            BsrAccountingProductOption._value_key(value)
            for value in templates.mapped("reporting_for")
            if value
        }
        reusable_set = len(template_keys) <= 1
        plan = []
        for option in options:
            scoped = (
                templates
                if reusable_set
                else templates.filtered(
                    lambda template: BsrAccountingProductOption._value_key(
                        template.reporting_for
                    )
                    == option.value_key
                )
            )
            plan.extend((template, option) for template in scoped)
        return plan

    def _visa_setup_for_product_option(self, product_option):
        self.ensure_one()
        return self.visa_setup_ids.filtered(
            lambda setup: product_option.value_key
            in setup.reporting_for_option_ids.mapped("value_key")
        )

    def _runtime_configuration(self, template, product_option=None, preflight=False):
        """Create the run-specific copy used for preflight and execution."""
        self.ensure_one()
        reporting_for = (
            product_option.name
            if product_option
            else template.reporting_for or template.product_type
        )
        setup = self._visa_setup_for_product_option(product_option) if product_option else False
        rollup_to = (
            ", ".join(setup.rollup_to_option_ids.filtered("active").sorted("name").mapped("name"))
            if setup
            else template.rollup_to
        )
        funds_xfer_entity = (
            ", ".join(
                setup.funds_xfer_entity_option_ids.filtered("active").sorted("name").mapped("name")
            )
            if setup
            else template.funds_xfer_entity
        )
        return template.sudo().with_context(
            tracking_disable=True,
            mail_notrack=True,
        ).copy(
            default={
                "name": (
                    _("Preflight — %s") % template.display_name
                    if preflight
                    else _("New")
                ),
                "template_id": template.id,
                "batch_run_id": self.id,
                "source_import_id": self.source_import_id.id,
                "production_date": self.production_date,
                "report_date": self.report_date,
                "posting_date": self.posting_date,
                # Product Type is the selected source Reporting For identity.
                "product_type": reporting_for,
                "reporting_for": reporting_for,
                "rollup_to": rollup_to,
                "funds_xfer_entity": funds_xfer_entity,
                "state": "draft",
                "validation_message": False,
                "matched_at": False,
                "matched_by_id": False,
            }
        )

    def _preflight_configuration_issues(self):
        """Return warnings/errors before creating any voucher reports."""
        self.ensure_one()
        issues = []
        covered_options = self.env["bsr.accounting.product.option"]
        for template, product_option in self._execution_plan():
            covered_options |= product_option
            runtime = self._runtime_configuration(
                template, product_option=product_option, preflight=True
            )
            try:
                match_result = runtime.with_context(allow_no_activity=True).action_match_source()
                if match_result.get("match_status") == "no_activity":
                    issues.append(
                        {
                            "severity": "warning",
                            "configuration": template,
                            "product_type": product_option.name,
                            "reason": runtime.validation_message,
                        }
                    )
                    continue
                self._validate_currency_setup(runtime)
                incomplete = runtime.line_ids.filtered(
                    lambda line: line.include_in_voucher
                    and line.configuration_status != "configured"
                )
                if incomplete:
                    samples = ", ".join(
                        filter(
                            None,
                            incomplete[:5].mapped("description_match"),
                        )
                    )
                    reason = _(
                        "%(count)s matched source rule(s) do not have a reusable GL mapping."
                    ) % {"count": len(incomplete)}
                    if samples:
                        reason += " " + _("Affected source descriptions: %s") % samples
                    raise UserError(reason)
            except UserError as error:
                issues.append(
                    {
                        "severity": "blocking",
                        "configuration": template,
                        "product_type": product_option.name,
                        "reason": str(error),
                    }
                )
            finally:
                runtime.unlink()
        for product_option in self.product_type_option_ids - covered_options:
            issues.append(
                {
                    "severity": "blocking",
                    "configuration": self.env["bsr.accounting.configuration"],
                    "product_type": product_option.name,
                    "reason": _(
                        "No selected configuration set is available for this Product Type."
                    ),
                }
            )
        return issues

    def _open_preflight_review(self, issues):
        self.ensure_one()
        source = self.source_import_id
        rows = []
        for issue in issues:
            template = issue["configuration"]
            template_source = template.source_import_id if template else False
            severity = issue["severity"]
            outcome = _("No Activity") if severity == "warning" else _("Blocking Issue")
            badge = "text-bg-warning" if severity == "warning" else "text-bg-danger"
            rows.append(
                "<tr>"
                f"<td><span class='badge {badge}'>{html.escape(outcome)}</span></td>"
                f"<td>{html.escape(issue['product_type'] or '')}</td>"
                f"<td><strong>{html.escape(template.display_name if template else _('None'))}</strong><br/>"
                f"<small>{html.escape(_('Database ID: #%s') % template.id) if template else ''}</small></td>"
                f"<td>{html.escape((template.source_report_codes or template.report_id_filter or '') if template else '')}</td>"
                f"<td>{html.escape(template_source.display_name or '') if template_source else ''}<br/>"
                f"<small>{html.escape(template_source.file_name or '') if template_source else ''}</small></td>"
                f"<td>{html.escape(issue['reason'] or '')}</td>"
                "</tr>"
            )
        source_details = _(
            "Processed VSS: %(reference)s (Database ID: #%(id)s) — File: %(file)s"
        ) % {
            "reference": source.display_name,
            "id": source.id,
            "file": source.file_name or _("Unknown"),
        }
        summary = (
            f"<div class='alert alert-info mb-3'><strong>{html.escape(source_details)}</strong></div>"
            "<div class='table-responsive'><table class='table table-sm table-bordered align-middle'>"
            "<thead><tr>"
            f"<th>{html.escape(_('Outcome'))}</th>"
            f"<th>{html.escape(_('Product Type / Reporting For'))}</th>"
            f"<th>{html.escape(_('Configuration'))}</th>"
            f"<th>{html.escape(_('Report IDs'))}</th>"
            f"<th>{html.escape(_('Original configuration file'))}</th>"
            f"<th>{html.escape(_('Exact reason'))}</th>"
            "</tr></thead><tbody>"
            + "".join(rows)
            + "</tbody></table></div>"
        )
        blocking_count = sum(issue["severity"] == "blocking" for issue in issues)
        wizard = self.env["bsr.accounting.run.review.wizard"].create(
            {
                "run_id": self.id,
                "issue_summary_html": summary,
                "warning_count": len(issues) - blocking_count,
                "blocking_count": blocking_count,
            }
        )
        return {
            "type": "ir.actions.act_window",
            "name": _("Voucher Pre-generation Review"),
            "res_model": "bsr.accounting.run.review.wizard",
            "view_mode": "form",
            "res_id": wizard.id,
            "target": "new",
        }

    def action_execute(self):
        self.ensure_one()
        if self.state != "draft" or self.generated_configuration_ids or self.report_ids:
            raise UserError(_("This run is already generated. Create a new run to preserve the audit trail."))
        self._prepare_visa_run_scope()
        self._validate_selected_configuration_currencies()
        self._synchronize_product_type_scope()
        if not self.product_type_option_ids:
            if self.product_type_selection_mode == "custom":
                raise UserError(
                    _("Select at least one Product Type when Product Type is set to Custom.")
                )
            raise UserError(
                _("No Product Types (Reporting For values) were found in the processed VSS file.")
            )
        if not self.configuration_ids:
            raise UserError(_("Select at least one completed configuration file."))
        self._apply_source_import_dates()
        invalid = self.configuration_ids.filtered(lambda item: item.state != "ready" or item.batch_run_id)
        if invalid:
            raise UserError(_("Only completed standalone configuration files can be executed."))
        if not self.env.context.get("skip_preflight_review"):
            issues = self._preflight_configuration_issues()
            if issues:
                return self._open_preflight_review(issues)
        operator = self.env.user
        allow_blocking_skip = bool(self.env.context.get("acknowledged_blocking_issues"))
        generated_reports = self.env["bsr.accounting.report"]
        for template, product_option in self._execution_plan():
            # Operators may execute approved templates without receiving
            # permission to alter the master configuration records.
            runtime = self._runtime_configuration(template, product_option=product_option)
            try:
                match_result = runtime.with_context(allow_no_activity=True).action_match_source()
                if match_result.get("match_status") == "no_activity":
                    continue
                self._validate_currency_setup(runtime)
                runtime.write({"matched_by_id": operator.id})
                incomplete = runtime.line_ids.filtered(
                    lambda line: line.include_in_voucher and line.configuration_status != "configured"
                )
                if incomplete:
                    raise UserError(
                        _("%(count)s source rule(s) do not have a reusable GL mapping.")
                        % {"count": len(incomplete)}
                    )
                runtime.action_validate()
            except UserError as error:
                if allow_blocking_skip:
                    runtime.line_ids.unlink()
                    runtime.write(
                        {
                            "state": "blocked",
                            "validation_message": _(
                                "Acknowledged blocking issue — configuration %(configuration)s "
                                "(Database ID: #%(configuration_id)s) was excluded from processed VSS "
                                "%(source)s (Database ID: #%(source_id)s), file %(file)s. Reason: %(reason)s"
                            )
                            % {
                                "configuration": template.display_name,
                                "configuration_id": template.id,
                                "source": self.source_import_id.display_name,
                                "source_id": self.source_import_id.id,
                                "file": self.source_import_id.file_name or _("Unknown"),
                                "reason": _("Product Type %(product)s: %(reason)s")
                                % {"product": product_option.name, "reason": str(error)},
                            },
                        }
                    )
                    continue
                raise UserError(
                    _(
                        "Configuration %(configuration)s cannot be applied to %(source)s.\n\n%(reason)s"
                    )
                    % {
                        "configuration": template.display_name,
                        "source": self.source_import_id.display_name,
                        "reason": _("Product Type %(product)s: %(reason)s")
                        % {"product": product_option.name, "reason": str(error)},
                    }
                ) from error
            # Report creation/generation is intentionally outside the
            # acknowledged-source-issue handler. An acknowledgment may skip a
            # configuration that cannot match this file, but it must never
            # suppress an unexpected voucher/report generation failure.
            report = self.env["bsr.accounting.report"].sudo().create(
                {"configuration_id": runtime.id, "batch_run_id": self.id}
            )
            report.action_generate()
            report.write({"generated_by_id": operator.id})
            generated_reports |= report

        if allow_blocking_skip and not generated_reports:
            raise UserError(
                _(
                    "No voucher can be generated because every selected configuration was either "
                    "blocked or had no activity. Resolve at least one blocking configuration first."
                )
            )

        self._generate_merged_output(generated_reports)
        product_groups = {}
        for report, mappings in self._merged_mapping_groups(generated_reports.sorted(
            lambda report: self._template_sort_key(report.configuration_id.template_id)
        )):
            product_type = report.configuration_id.reporting_for or _("Unspecified")
            product_groups.setdefault(product_type, []).append((report, mappings))
        products_balanced = all(
            self._voucher_section_data(groups)["difference"] < 0.005
            for groups in product_groups.values()
        )
        self.write(
            {
                "state": (
                    "done"
                    if products_balanced and not self.blocked_configuration_count
                    else "exception"
                ),
                "generated_at": fields.Datetime.now(),
                "generated_by_id": self.env.user.id,
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{self._name}/{self.id}/output_file/{self.output_filename}?download=true",
            "target": "self",
        }

    def _voucher_section_data(self, mapping_groups):
        """Build a voucher dataset with BDT FX balancing and USD control totals."""
        self.ensure_one()
        lines = []
        additional_fields = set()
        currency_split = any(
            report.configuration_id.show_currency_columns
            or any(
                _mapping_amount_currency(mapping, mapping.amount_source) == "USD"
                for mapping in mappings
            )
            for report, mappings in mapping_groups
        )
        for _report, mappings in mapping_groups:
            for mapping in mappings:
                amount_usd, amount_bdt, _rate, remaining = _mapping_currency_values(mapping, self)
                if not currency_split:
                    amount_bdt = mapping.matched_amount
                additional_fields.update(remaining)
                lines.append(
                    {
                        "mapping": mapping,
                        "side": mapping.posting_indicator,
                        "amount_usd": amount_usd,
                        "amount_bdt": amount_bdt,
                        "additional": remaining,
                    }
                )
        additional_fields = sorted(
            additional_fields,
            key=lambda field: (
                AMOUNT_SOURCE_ORDER.index(field)
                if field in AMOUNT_SOURCE_ORDER
                else len(AMOUNT_SOURCE_ORDER),
                field,
            ),
        )
        raw_debit_usd = sum(line["amount_usd"] for line in lines if line["side"] == "DR")
        raw_credit_usd = sum(line["amount_usd"] for line in lines if line["side"] == "CR")
        raw_debit_bdt = sum(line["amount_bdt"] for line in lines if line["side"] == "DR")
        raw_credit_bdt = sum(line["amount_bdt"] for line in lines if line["side"] == "CR")
        debit_usd, credit_usd = raw_debit_usd, raw_credit_usd
        debit_bdt, credit_bdt = raw_debit_bdt, raw_credit_bdt
        adjustments = (
            _currency_balance_adjustments(debit_usd, credit_usd, debit_bdt, credit_bdt)
            if currency_split
            else []
        )
        for adjustment in adjustments:
            if adjustment["side"] == "DR":
                debit_usd += adjustment["amount_usd"]
                debit_bdt += adjustment["amount_bdt"]
            else:
                credit_usd += adjustment["amount_usd"]
                credit_bdt += adjustment["amount_bdt"]
        usd_difference = abs(debit_usd - credit_usd)
        bdt_difference = abs(debit_bdt - credit_bdt)
        return {
            "lines": lines,
            "currency_split": currency_split,
            "additional_fields": additional_fields,
            "adjustments": adjustments,
            "adjustment": adjustments[0] if adjustments else False,
            "raw_debit_usd": raw_debit_usd,
            "raw_credit_usd": raw_credit_usd,
            "raw_debit_bdt": raw_debit_bdt,
            "raw_credit_bdt": raw_credit_bdt,
            "debit_usd": debit_usd,
            "credit_usd": credit_usd,
            "debit_bdt": debit_bdt,
            "credit_bdt": credit_bdt,
            "usd_difference": usd_difference,
            "bdt_difference": bdt_difference,
            "difference": max(usd_difference, bdt_difference),
            "provisional": bool(self.blocked_configuration_count),
        }

    def _write_voucher_table(self, sheet, row, data):
        """Render ordered source postings and one product-level FX adjustment."""
        navy = "123A5A"
        pale_green = "E2F0D9"
        alternate = "F7FAFC"
        amount_format = '#,##0.00;[Red](#,##0.00)'
        thin = Side(style="thin", color="808080")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        currency_split = data["currency_split"]
        fixed_end = 10 if currency_split else 8
        # Source amount fields are already represented by debit/credit values.
        # They remain available on the audit sheet, not duplicated here.
        table_end = fixed_end
        headers = (
            ["GL Number", "GL Name", "Debit Amount (USD)", "Debit Amount (BDT)", "Debit Remarks",
             "GL Number", "GL Name", "Credit Amount (USD)", "Credit Amount (BDT)", "Credit Remarks"]
            if currency_split else
            ["GL Number", "GL Name", "Debit Amount", "Debit Remarks",
             "GL Number", "GL Name", "Credit Amount", "Credit Remarks"]
        )
        header_row = row
        for column, label in enumerate(headers, 1):
            cell = sheet.cell(row, column, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 34
        row += 1

        for index, item in enumerate(data["lines"]):
            mapping = item["mapping"]
            if item["side"] == "DR":
                values = (
                    [mapping.gl_number, mapping.gl_name, item["amount_usd"], item["amount_bdt"],
                     mapping.debit_remarks or mapping.remarks or ""]
                    if currency_split else
                    [mapping.gl_number, mapping.gl_name, item["amount_bdt"],
                     mapping.debit_remarks or mapping.remarks or ""]
                )
                start_column = 1
            else:
                values = (
                    [mapping.gl_number, mapping.gl_name, item["amount_usd"], item["amount_bdt"],
                     mapping.credit_remarks or mapping.remarks or ""]
                    if currency_split else
                    [mapping.gl_number, mapping.gl_name, item["amount_bdt"],
                     mapping.credit_remarks or mapping.remarks or ""]
                )
                start_column = 6 if currency_split else 5
            for offset, value in enumerate(values):
                sheet.cell(row, start_column + offset, value)
            sheet.cell(row, start_column).number_format = "@"
            amount_columns = (3, 4, 8, 9) if currency_split else (3, 7)
            for column in amount_columns:
                sheet.cell(row, column).number_format = amount_format
            for column in range(1, table_end + 1):
                cell = sheet.cell(row, column)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="right" if column in amount_columns else "left",
                    vertical="center",
                    wrap_text=column in ((2, 5, 7, 10) if currency_split else (2, 4, 6, 8)),
                )
                if index % 2:
                    cell.fill = PatternFill("solid", fgColor=alternate)
            debit_remark = mapping.debit_remarks or mapping.remarks or ""
            credit_remark = mapping.credit_remarks or mapping.remarks or ""
            sheet.row_dimensions[row].height = _voucher_row_height(
                (mapping.gl_name, 32),
                (debit_remark if item["side"] == "DR" else credit_remark, 38),
            )
            row += 1

        for adjustment in data["adjustments"]:
            start_column = 1 if adjustment["side"] == "DR" else (6 if currency_split else 5)
            amount_column = 4 if adjustment["side"] == "DR" else 9
            remarks_column = 5 if adjustment["side"] == "DR" else 10
            sheet.cell(row, start_column, self.currency_conversion_gl_number)
            sheet.cell(row, start_column + 1, adjustment["label"])
            sheet.cell(row, start_column + 2, adjustment["amount_usd"])
            sheet.cell(row, amount_column, adjustment["amount_bdt"])
            sheet.cell(row, remarks_column, adjustment["label"])
            sheet.cell(row, start_column).number_format = "@"
            for column in (3, 4, 8, 9):
                sheet.cell(row, column).number_format = amount_format
            for column in range(1, table_end + 1):
                cell = sheet.cell(row, column)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="right" if column in (3, 4, 8, 9) else "left",
                    vertical="center",
                    wrap_text=column in (2, 5, 7, 10),
                )
            sheet.row_dimensions[row].height = _voucher_row_height(
                (adjustment["label"], 32), (adjustment["label"], 38)
            )
            row += 1

        if currency_split:
            totals = {
                2: "Total", 3: data["debit_usd"], 4: data["debit_bdt"],
                7: "Total", 8: data["credit_usd"], 9: data["credit_bdt"],
            }
            amount_columns = (3, 4, 8, 9)
        else:
            totals = {2: "Total", 3: data["debit_bdt"], 6: "Total", 7: data["credit_bdt"]}
            amount_columns = (3, 7)
        for column, value in totals.items():
            sheet.cell(row, column, value).font = Font(bold=True)
        for column in range(1, table_end + 1):
            cell = sheet.cell(row, column)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=pale_green)
            cell.alignment = Alignment(
                horizontal="right" if column in amount_columns else "left",
                vertical="center",
            )
        for column in amount_columns:
            sheet.cell(row, column).number_format = amount_format
        sheet.row_dimensions[row].height = 23
        row += 1
        if currency_split:
            diff_label_col, status_col = 7, 10
            difference_columns = ((8, data["usd_difference"]), (9, data["bdt_difference"]))
        else:
            diff_label_col, status_col = 6, 8
            difference_columns = ((7, data["difference"]),)
        difference = data["difference"]
        sheet.cell(row, diff_label_col, _("Difference")).font = Font(bold=True)
        for column, value in difference_columns:
            sheet.cell(row, column, value)
            sheet.cell(row, column).number_format = amount_format
        if difference < 0.005 and data.get("provisional"):
            status = _("BALANCED — ACKNOWLEDGED EXCEPTIONS")
            color = "C65911"
        else:
            status = _("BALANCED") if difference < 0.005 else _("MISMATCH FLAG")
            color = "008000" if difference < 0.005 else "C00000"
        sheet.cell(row, status_col, status)
        for column, _value in difference_columns:
            sheet.cell(row, column).font = Font(bold=True, color=color)
        sheet.cell(row, status_col).font = Font(bold=True, color=color)
        sheet.cell(row, status_col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.row_dimensions[row].height = _voucher_row_height((status, 30))
        return row, header_row, table_end

    def _write_product_section(self, sheet, row, product_type, mapping_groups):
        """Append one combined voucher for a unique product type."""
        data = self._voucher_section_data(mapping_groups)
        table_end = 10 if data["currency_split"] else 8
        navy = "123A5A"
        yellow = "FFF200"
        row += 4
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        cell = sheet.cell(row, 1, _("Product Voucher — %s") % product_type)
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[row].height = 25
        row += 2
        reports = self.env["bsr.accounting.report"]
        for report, _mappings in mapping_groups:
            reports |= report
        owners = " | ".join(dict.fromkeys(filter(None, reports.mapped("configuration_id.business_owner"))))
        reporting_for = " | ".join(dict.fromkeys(filter(None, reports.mapped("configuration_id.reporting_for"))))
        for title in (
            self.company_id.name,
            owners,
            reporting_for,
            _("Settlement (VSS) Date: %s") % fields.Date.to_string(self.report_date),
            _("Posting Date: %s") % fields.Date.to_string(self.posting_date),
        ):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
            title_cell = sheet.cell(row, 1, title)
            title_cell.font = Font(bold=True, color=navy)
            title_cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            sheet.row_dimensions[row].height = _voucher_row_height((title, 100))
            row += 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        product_cell = sheet.cell(row, 1, product_type)
        product_cell.fill = PatternFill("solid", fgColor=yellow)
        product_cell.font = Font(bold=True)
        product_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 23
        row += 2
        row, _header_row, table_end = self._write_voucher_table(sheet, row, data)
        row += 3
        sheet.cell(row, 1, _("Maker")).font = Font(bold=True)
        sheet.cell(row, 2, self.generated_by_id.name or self.env.user.name)
        checker_names = " | ".join(dict.fromkeys(filter(None, reports.mapped("checked_by_id.name"))))
        checker_label_col, checker_value_col = ((7, 8) if data["currency_split"] else (6, 7))
        sheet.cell(row, checker_label_col, _("Checker")).font = Font(bold=True)
        sheet.cell(row, checker_value_col, checker_names)
        return row, table_end

    def _write_product_status_section(self, sheet, row, product_type):
        """Keep a selected Product Type visible when it produced no voucher."""
        table_end = 10
        navy = "123A5A"
        pale = "FFF2CC"
        row += 4
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        title = sheet.cell(row, 1, _("Product Voucher — %s") % product_type)
        title.font = Font(bold=True, color="FFFFFF", size=12)
        title.fill = PatternFill("solid", fgColor=navy)
        title.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[row].height = 25
        row += 2
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        message = sheet.cell(
            row,
            1,
            _(
                "No voucher was generated for this Product Type. All assigned configurations "
                "were No Activity or were skipped after acknowledged blocking issues."
            ),
        )
        message.fill = PatternFill("solid", fgColor=pale)
        message.font = Font(bold=True, color="9C6500")
        message.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 38
        return row, table_end

    def _generate_merged_output(self, reports=None):
        self.ensure_one()
        if Workbook is None:
            raise UserError(_("The Python package 'openpyxl' is required to generate the merged report."))
        reports = reports or self.report_ids
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Merged Voucher Report"
        sheet.sheet_view.showGridLines = False
        ordered_reports = reports.sorted(lambda report: self._template_sort_key(report.configuration_id.template_id))
        mapping_groups = self._merged_mapping_groups(ordered_reports)
        merged_data = self._voucher_section_data(mapping_groups)
        currency_split = merged_data["currency_split"]
        fixed_end = 10 if currency_split else 8
        table_end = fixed_end
        navy = "123A5A"
        yellow = "FFF200"
        pale = "FFF2CC"
        thin = Side(style="thin", color="808080")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        widths = (
            {"A": 20, "B": 32, "C": 18, "D": 18, "E": 38,
             "F": 20, "G": 32, "H": 18, "I": 18, "J": 38}
            if currency_split else
            {"A": 20, "B": 32, "C": 18, "D": 38,
             "E": 20, "F": 32, "G": 18, "H": 38}
        )
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        summary = [
            ("Report ID", self.requested_report_ids),
            ("Product Type", self.product_type),
            ("Processed VSS File", self.source_import_id.display_name),
            ("PROC Date", fields.Date.to_string(self.production_date)),
            ("Report Date", fields.Date.to_string(self.report_date)),
            ("Posting Date", fields.Date.to_string(self.posting_date)),
            ("Configuration Files", len(self.configuration_ids)),
            ("Generated Reports", len(ordered_reports)),
            ("No Activity", self.no_activity_count),
            ("Acknowledged Blocking Issues", self.blocked_configuration_count),
            ("Currency Conversion GL", self.currency_conversion_gl_number or ""),
            ("USD Buy Rate", self.usd_buy_rate),
            ("USD Sale Rate", self.usd_sale_rate),
            ("USD Mid Rate", self.usd_mid_rate),
        ]
        row = 1
        for label, value in summary:
            label_cell = sheet.cell(row, 1, label)
            value_cell = sheet.cell(row, 2, value)
            label_cell.font = Font(bold=True, color=navy)
            label_cell.fill = PatternFill("solid", fgColor=pale)
            value_cell.fill = PatternFill("solid", fgColor=pale)
            label_cell.border = border
            value_cell.border = border
            value_cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 19
            row += 1

        business_owners = " | ".join(
            dict.fromkeys(filter(None, ordered_reports.mapped("configuration_id.business_owner")))
        )
        reporting_for = " | ".join(
            dict.fromkeys(filter(None, ordered_reports.mapped("configuration_id.reporting_for")))
        )
        row += 1
        for title in (
            self.company_id.name,
            business_owners or _("Merged Voucher"),
            reporting_for or self.requested_report_ids,
            _("Settlement (VSS) Date: %s") % fields.Date.to_string(self.report_date),
            _("Posting Date: %s") % fields.Date.to_string(self.posting_date),
        ):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
            cell = sheet.cell(row, 1, title)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(bold=True, color=navy)
            row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        product_cell = sheet.cell(row, 1, self.product_type)
        product_cell.fill = PatternFill("solid", fgColor=yellow)
        product_cell.font = Font(bold=True)
        product_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 23
        row += 2

        row, header_row, table_end = self._write_voucher_table(sheet, row, merged_data)
        row += 3
        sheet.cell(row, 1, "Maker").font = Font(bold=True)
        sheet.cell(row, 2, self.generated_by_id.name or self.env.user.name)
        checker_names = " | ".join(dict.fromkeys(filter(None, reports.mapped("checked_by_id.name"))))
        checker_label_col, checker_value_col = ((7, 8) if currency_split else (6, 7))
        sheet.cell(row, checker_label_col, "Checker").font = Font(bold=True)
        sheet.cell(row, checker_value_col, checker_names)
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True

        details = workbook.create_sheet("Merged Source Details")
        grouped_details = []
        amount_fields = set()
        for report, mappings in mapping_groups:
            groups, grouped_fields = _group_source_details(mappings)
            grouped_details.extend((report, group) for group in groups)
            amount_fields.update(grouped_fields)
        amount_fields = sorted(
            amount_fields,
            key=lambda field: (AMOUNT_SOURCE_ORDER.index(field) if field in AMOUNT_SOURCE_ORDER else len(AMOUNT_SOURCE_ORDER), field),
        )
        headers = [
            "Configuration", "Voucher", "Report ID", "Source Level", "Business Mode",
            "Transaction Category", "Transaction Type", "Source Description", "Source Currency",
            "Source CR/DB", "Posting DR/CR",
        ]
        headers.extend(AMOUNT_SOURCE_LABELS.get(field, field) for field in amount_fields or ("matched_amount",))
        headers.extend([
            "GL Number", "GL Name", "Source Row Count", "Parsed Row IDs",
            "Debit Remarks", "Credit Remarks",
            "Execution Status", "Execution Message",
        ])
        for column, label in enumerate(headers, 1):
            detail_cell = details.cell(1, column, label)
            detail_cell.font = Font(bold=True, color="FFFFFF")
            detail_cell.fill = PatternFill("solid", fgColor=navy)
        detail_row = 2
        amount_start = 12
        for report, group in grouped_details:
            line = group["line"]
            values = [
                line.configuration_id.name, report.name, line.report_code,
                line.record_type, line.business_mode, line.transaction_category,
                line.transaction_type, line.description_match, line.source_currency,
                " / ".join(sorted(group["source_directions"])),
                " / ".join(sorted(group["posting_indicators"])),
            ]
            values.extend(group["amounts"].get(field, 0.0) for field in amount_fields or ("matched_amount",))
            values.extend([
                " / ".join(sorted(group["gl_numbers"])),
                " / ".join(sorted(group["gl_names"])),
                group["row_count"],
                ",".join(dict.fromkeys(group["record_ids"])),
                " | ".join(sorted(group["debit_remarks"])),
                " | ".join(sorted(group["credit_remarks"])),
                _("Generated"),
                "",
            ])
            for column, value in enumerate(values, 1):
                details.cell(detail_row, column, value)
            for column in range(amount_start, amount_start + (len(amount_fields) or 1)):
                details.cell(detail_row, column).number_format = '#,##0.00;[Red](#,##0.00)'
            detail_row += 1
        for runtime in self.generated_configuration_ids.filtered(
            lambda config: config.state in {"no_activity", "blocked"}
        ).sorted(self._template_sort_key):
            execution_status = (
                _("No Activity")
                if runtime.state == "no_activity"
                else _("Skipped — Acknowledged Blocking Issue")
            )
            values = [
                runtime.template_id.display_name or runtime.display_name,
                "",
                runtime.source_report_codes,
            ]
            values.extend([""] * (len(headers) - len(values) - 2))
            values.extend([execution_status, runtime.validation_message or ""])
            for column, value in enumerate(values, 1):
                details.cell(detail_row, column, value)
            detail_row += 1
        details.freeze_panes = "A2"
        details.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(detail_row - 1, 2)}"
        widths = [24, 24, 14, 14, 18, 24, 24, 32, 17, 15, 16]
        widths.extend([18] * (len(amount_fields) or 1))
        widths.extend([16, 28, 17, 28, 28, 28, 18, 60])
        for column, width in enumerate(widths, 1):
            details.column_dimensions[details.cell(1, column).column_letter].width = width

        # For multi-product runs, append one combined voucher per selected
        # Reporting For value after the merged voucher. Keep products with no
        # generated report visible instead of silently omitting them.
        product_groups = {}
        for report, mappings in mapping_groups:
            product_key = BsrAccountingProductOption._value_key(
                report.configuration_id.reporting_for or _("Unspecified")
            )
            product_groups.setdefault(product_key, []).append((report, mappings))
        max_table_end = table_end
        if len(self.product_type_option_ids) > 1:
            for product_option in self.product_type_option_ids.sorted("name"):
                product_mapping_groups = product_groups.get(product_option.value_key)
                if product_mapping_groups:
                    row, product_table_end = self._write_product_section(
                        sheet, row, product_option.name, product_mapping_groups
                    )
                else:
                    row, product_table_end = self._write_product_status_section(
                        sheet, row, product_option.name
                    )
                max_table_end = max(max_table_end, product_table_end)

        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.print_title_rows = f"{header_row}:{header_row}"
        sheet.print_area = f"A1:{get_column_letter(max_table_end)}{row}"

        output = io.BytesIO()
        workbook.save(output)
        self.write(
            {
                "output_file": base64.b64encode(output.getvalue()),
                "output_filename": f"{self.name.replace('/', '-')}-merged.xlsx",
            }
        )

    def action_download(self):
        self.ensure_one()
        if not self.output_file:
            raise UserError(_("Execute the configuration run first."))
        # Rebuild on every download so completed runs receive the current
        # professional single-table layout without rerunning source matching.
        self._generate_merged_output()
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{self._name}/{self.id}/output_file/{self.output_filename}?download=true",
            "target": "self",
        }


class BsrAccountingRunReviewWizard(models.TransientModel):
    _name = "bsr.accounting.run.review.wizard"
    _description = "Voucher Run Pre-generation Review"

    run_id = fields.Many2one(
        "bsr.accounting.run",
        required=True,
        readonly=True,
        ondelete="cascade",
    )
    issue_summary_html = fields.Html(readonly=True, sanitize=True)
    warning_count = fields.Integer(readonly=True)
    blocking_count = fields.Integer(readonly=True)
    acknowledged = fields.Boolean(
        string="Acknowledgment",
        help=(
            "Confirm that configurations with blocking issues may be excluded while valid "
            "configurations continue to voucher generation."
        ),
    )
    has_blocking_issues = fields.Boolean(compute="_compute_has_blocking_issues")

    @api.depends("blocking_count")
    def _compute_has_blocking_issues(self):
        for wizard in self:
            wizard.has_blocking_issues = bool(wizard.blocking_count)

    def action_continue(self):
        self.ensure_one()
        if not self.acknowledged:
            raise UserError(_("Acknowledge the listed issues before generating the voucher."))
        if self.has_blocking_issues:
            valid_count = (
                len(self.run_id._execution_plan())
                - self.warning_count
                - self.blocking_count
            )
            if valid_count <= 0:
                raise UserError(
                    _(
                        "No valid configuration remains for voucher generation. "
                        "Resolve at least one blocking issue first."
                    )
                )
        now = fields.Datetime.now()
        self.run_id.write(
            {
                "issues_acknowledged": True,
                "issues_acknowledged_at": now,
                "issues_acknowledged_by_id": self.env.user.id,
            }
        )
        self.run_id.message_post(
            body=_(
                "Pre-generation issues acknowledged by %(user)s. "
                "No Activity: %(warnings)s; blocking configurations excluded: %(blocking)s."
            )
            % {
                "user": self.env.user.display_name,
                "warnings": self.warning_count,
                "blocking": self.blocking_count,
            }
        )
        return self.run_id.with_context(
            skip_preflight_review=True,
            acknowledged_blocking_issues=bool(self.blocking_count),
        ).action_execute()

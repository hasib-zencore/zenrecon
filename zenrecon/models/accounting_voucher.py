# -*- coding: utf-8 -*-
import base64
import io

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None


AMOUNT_SOURCE_LABELS = {
    "auto": "Automatic (Legacy)",
    "total_amount": "Net Settlement Amount",
    "debit_amount": "VSS-110 Debit Amount",
    "credit_amount": "VSS-110 Credit Amount",
    "clearing_amount": "Clearing Amount",
    "interchange_amount": "Interchange Amount",
    "interchange_credit": "Interchange Credit",
    "interchange_debit": "Interchange Debit",
    "fee_credit": "Reimbursement Fee Credit",
    "fee_debit": "Reimbursement Fee Debit",
    "visa_charge_credit": "Visa Charge Credit",
    "visa_charge_debit": "Visa Charge Debit",
    "settlement_interchange_amount": "Settlement Interchange Amount",
    "settlement_conversion_fee": "Settlement Conversion Fee",
    "clearing_interchange_amount": "Clearing Interchange Amount",
    "clearing_conversion_fee": "Clearing Conversion Fee",
    "optional_issuer_fee": "Optional Issuer Fee",
}
AMOUNT_SOURCE_ORDER = tuple(AMOUNT_SOURCE_LABELS)

_CLEARING_CURRENCY_FIELDS = {
    "VSS-120": {"clearing_amount"},
    "VSS-210": {
        "clearing_interchange_amount",
        "clearing_conversion_fee",
        "optional_issuer_fee",
    },
}

_CURRENCY_PAIR_FIELDS = {
    "VSS-210": {
        "settlement_interchange_amount": "clearing_interchange_amount",
        "clearing_interchange_amount": "settlement_interchange_amount",
        "settlement_conversion_fee": "clearing_conversion_fee",
        "clearing_conversion_fee": "settlement_conversion_fee",
    },
}


def _mapping_amount_currency(mapping, amount_field):
    """Return the source currency of one selected VSS amount field.

    Runtime mappings contain the currency parsed from the actual VSS row.  It
    is more authoritative than a reusable configuration's optional currency
    filter, which may legitimately be blank or may belong to an older file.
    """
    config = mapping.configuration_id
    if amount_field == mapping.amount_source:
        source_currency = (mapping.source_currency or "").strip().upper()
        if source_currency:
            return source_currency
    if amount_field in _CLEARING_CURRENCY_FIELDS.get(mapping.report_code, set()):
        return (
            mapping.clearing_currency
            or config.clearing_currency
            or mapping.source_currency
            or ""
        ).strip().upper()
    return (mapping.settlement_currency or config.settlement_currency or "").strip().upper()


def _paired_amount_field(mapping):
    """Find the opposite-currency amount that belongs on the same GL row."""
    extras = mapping.amount_columns_json or {}
    candidate = False
    if mapping.report_code == "VSS-120":
        if mapping.amount_source == "clearing_amount":
            candidate = "interchange_debit" if mapping.posting_indicator == "DR" else "interchange_credit"
        elif mapping.amount_source in {"interchange_credit", "interchange_debit"}:
            candidate = "clearing_amount"
    if not candidate:
        candidate = _CURRENCY_PAIR_FIELDS.get(mapping.report_code, {}).get(mapping.amount_source)
    if candidate not in extras:
        return False
    primary_currency = _mapping_amount_currency(mapping, mapping.amount_source)
    pair_currency = _mapping_amount_currency(mapping, candidate)
    return candidate if primary_currency and pair_currency and primary_currency != pair_currency else False


def _mapping_currency_values(mapping, run=None):
    """Return USD/BDT values, rate and informational fields for one GL row.

    If the source supplies both currencies, both original values are retained.
    Conversion is applied only when the accounting row supplies USD alone.
    """
    config = mapping.configuration_id
    # A standalone single-currency voucher keeps its original amount. During
    # a batch run, USD-only product rules still participate in the run-level
    # USD/BDT conversion and balancing.
    if not run and not config.show_usd_amount:
        return 0.0, mapping.matched_amount, 0.0, mapping.amount_columns_json or {}

    primary_currency = _mapping_amount_currency(mapping, mapping.amount_source)
    pair_field = _paired_amount_field(mapping)
    pair_currency = _mapping_amount_currency(mapping, pair_field) if pair_field else False
    amount_usd = 0.0
    amount_bdt = 0.0
    conversion_rate = 0.0
    if primary_currency == "USD":
        amount_usd = mapping.matched_amount
    else:
        amount_bdt = mapping.matched_amount
    if pair_field:
        pair_amount = (mapping.amount_columns_json or {}).get(pair_field, 0.0) or 0.0
        if pair_currency == "USD":
            amount_usd = pair_amount
        elif pair_currency == "BDT":
            amount_bdt = pair_amount
    if amount_usd and not amount_bdt:
        conversion_rate = run._conversion_rate_for(mapping.conversion_type) if run else 0.0
        amount_bdt = amount_usd * conversion_rate
    remaining = {
        field: amount
        for field, amount in (mapping.amount_columns_json or {}).items()
        if field != pair_field
    }
    return amount_usd, amount_bdt, conversion_rate, remaining


def _currency_balance_adjustments(debit_usd, credit_usd, debit_bdt, credit_bdt):
    """Return the BDT-only currency conversion gain/loss adjustment.

    USD is the source-currency control ledger and must never be forced into
    balance through the conversion GL.  Any USD variance therefore remains
    visible as a genuine voucher mismatch.  Only the converted BDT ledger is
    adjusted for the exchange-rate gain or loss.
    """
    del debit_usd, credit_usd  # Kept in the signature for existing callers.
    residual = debit_bdt - credit_bdt
    if abs(residual) < 0.005:
        return []
    side = "CR" if residual > 0 else "DR"
    gain_loss = "gain" if side == "CR" else "loss"
    return [
        {
            "side": side,
            "label": "Currency conversion BDT %s" % gain_loss,
            "amount_usd": 0.0,
            "amount_bdt": abs(residual),
        }
    ]


def _additional_amount_fields(lines):
    """Return the extra selected amount fields in their configured display order."""
    selected = {
        field
        for line in lines
        for field in (line.amount_columns_json or {})
    }
    return sorted(
        selected,
        key=lambda field: (
            AMOUNT_SOURCE_ORDER.index(field)
            if field in AMOUNT_SOURCE_ORDER
            else len(AMOUNT_SOURCE_ORDER),
            field,
        ),
    )


def _voucher_row_height(*text_width_pairs):
    """Estimate a readable Excel row height for wrapped voucher text."""
    wrapped_lines = 1
    for value, character_width in text_width_pairs:
        text = str(value or "")
        line_count = sum(
            max(1, (len(part) + character_width - 1) // character_width)
            for part in (text.splitlines() or [""])
        )
        wrapped_lines = max(wrapped_lines, line_count)
    return min(72, max(24, 15 * wrapped_lines + 4))


def _group_source_details(lines):
    """Combine selected amount sources into columns for the audit sheet only."""
    def value(line, *names):
        for name in names:
            if name in line._fields:
                return line[name]
        return False

    groups = {}
    for line in lines.sorted("sequence"):
        key = (
            value(line, "source_report_code", "report_code"),
            value(line, "source_record_type", "record_type"),
            value(line, "source_business_mode", "business_mode"),
            value(line, "source_transaction_category", "transaction_category"),
            value(line, "source_transaction_type", "transaction_type"),
            value(line, "description", "description_match"),
            value(line, "source_currency"),
        )
        group = groups.setdefault(
            key,
            {
                "line": line,
                "amounts": {},
                "source_directions": set(),
                "posting_indicators": set(),
                "gl_numbers": set(),
                "gl_names": set(),
                "debit_remarks": set(),
                "credit_remarks": set(),
                "row_count": 0,
                "record_ids": [],
            },
        )
        amount_key = value(line, "source_amount_field", "amount_source") or "matched_amount"
        amount = value(line, "amount", "matched_amount") or 0.0
        group["amounts"][amount_key] = group["amounts"].get(amount_key, 0.0) + amount
        for field, extra_amount in (value(line, "amount_columns_json") or {}).items():
            group["amounts"][field] = group["amounts"].get(field, 0.0) + (extra_amount or 0.0)
        source_direction = value(line, "source_direction")
        posting_indicator = value(line, "posting_indicator")
        gl_number = value(line, "gl_number")
        gl_name = value(line, "gl_name")
        if source_direction:
            group["source_directions"].add(source_direction)
        if posting_indicator:
            group["posting_indicators"].add(posting_indicator)
        if gl_number:
            group["gl_numbers"].add(gl_number)
        if gl_name:
            group["gl_names"].add(gl_name)
        group["row_count"] += value(line, "source_row_count") or 0
        group["record_ids"].extend(filter(None, (value(line, "source_record_ids") or "").split(",")))
        remarks = value(line, "remarks")
        debit_remarks = value(line, "debit_remarks") or (
            remarks if source_direction in {"DB", "DR"} else ""
        )
        credit_remarks = value(line, "credit_remarks") or (
            remarks if source_direction == "CR" else ""
        )
        if debit_remarks and str(debit_remarks).strip():
            group["debit_remarks"].add(str(debit_remarks).strip())
        if credit_remarks and str(credit_remarks).strip():
            group["credit_remarks"].add(str(credit_remarks).strip())
    amount_fields = sorted(
        {field for group in groups.values() for field in group["amounts"]},
        key=lambda field: (AMOUNT_SOURCE_ORDER.index(field) if field in AMOUNT_SOURCE_ORDER else len(AMOUNT_SOURCE_ORDER), field),
    )
    return list(groups.values()), amount_fields


class BsrAccountingReport(models.Model):
    _name = "bsr.accounting.report"
    _description = "Generated Bank Reconciliation Voucher Report"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "create_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(default=lambda self: _("New"), readonly=True, copy=False, tracking=True)
    configuration_id = fields.Many2one(
        "bsr.accounting.configuration", required=True, index=True, ondelete="restrict", check_company=True
    )
    batch_run_id = fields.Many2one(
        "bsr.accounting.run", string="Configuration Run", readonly=True, copy=False, index=True, ondelete="cascade"
    )
    company_id = fields.Many2one(related="configuration_id.company_id", store=True, index=True)
    source_import_id = fields.Many2one(related="configuration_id.source_import_id", store=True)
    report_code = fields.Selection(related="configuration_id.report_code", store=True)
    source_report_codes = fields.Char(related="configuration_id.source_report_codes", store=True)
    product_type = fields.Char(related="configuration_id.product_type", store=True)
    production_date = fields.Date(related="configuration_id.production_date", store=True)
    report_date = fields.Date(related="configuration_id.report_date", store=True)
    posting_date = fields.Date(related="configuration_id.posting_date", store=True)
    voucher_ids = fields.One2many("bsr.accounting.voucher", "report_id", string="Voucher Details")
    voucher_id = fields.Many2one("bsr.accounting.voucher", readonly=True, copy=False)
    state = fields.Selection(
        [
            ("draft", "Draft"),
            ("balanced", "Balanced"),
            ("exception", "Exception"),
            ("checked", "Checked"),
            ("approved", "Approved"),
            ("cancelled", "Cancelled"),
        ],
        default="draft",
        required=True,
        tracking=True,
        index=True,
    )
    debit_total = fields.Float(related="voucher_id.debit_total", digits=(20, 2))
    credit_total = fields.Float(related="voucher_id.credit_total", digits=(20, 2))
    difference = fields.Float(related="voucher_id.difference", digits=(20, 2))
    issue_count = fields.Integer(related="voucher_id.issue_count")
    generated_at = fields.Datetime(readonly=True)
    generated_by_id = fields.Many2one("res.users", readonly=True, string="Maker")
    checked_at = fields.Datetime(readonly=True)
    checked_by_id = fields.Many2one("res.users", readonly=True, string="Checker")
    approved_at = fields.Datetime(readonly=True)
    approved_by_id = fields.Many2one("res.users", readonly=True)
    output_file = fields.Binary(attachment=True, readonly=True, copy=False)
    output_filename = fields.Char(readonly=True, copy=False)
    output_mimetype = fields.Char(readonly=True, copy=False)

    @api.model_create_multi
    def create(self, vals_list):
        for values in vals_list:
            if not values.get("name") or values.get("name") == _("New"):
                values["name"] = self.env["ir.sequence"].next_by_code("bsr.accounting.report") or _("New")
        return super().create(vals_list)

    def action_generate(self):
        for report in self:
            if report.configuration_id.state != "ready":
                raise UserError(_("The configuration must be complete before a voucher can be generated."))
            if report.voucher_id:
                report.voucher_id.unlink()
            voucher = self.env["bsr.accounting.voucher"].create({"report_id": report.id})
            report.voucher_id = voucher
            voucher._build_from_configuration()
            report.write(
                {
                    "state": "balanced" if voucher.state == "balanced" else "exception",
                    "generated_at": fields.Datetime.now(),
                    "generated_by_id": self.env.user.id,
                    "checked_at": False,
                    "checked_by_id": False,
                    "approved_at": False,
                    "approved_by_id": False,
                }
            )
            report._generate_output_file()
        return True

    def _generate_output_file(self):
        self.ensure_one()
        if Workbook is None:
            raise UserError(_("The Python package 'openpyxl' is required to generate the voucher report."))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Voucher Report"
        self._write_sheet(sheet)
        self._write_source_details(workbook.create_sheet("Source Details"))
        output = io.BytesIO()
        workbook.save(output)
        filename = f"{self.name.replace('/', '-')}.xlsx"
        self.write(
            {
                "output_file": base64.b64encode(output.getvalue()),
                "output_filename": filename,
                "output_mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            }
        )

    def _write_sheet(self, sheet, start_row=1, configure_sheet=True):
        """Write one voucher section and return its final row.

        ``start_row`` lets a batch workbook place configuration vouchers below
        the merged voucher without creating one worksheet per configuration.
        """
        config = self.configuration_id
        voucher = self.voucher_id
        currency_split = bool(voucher.show_currency_columns)
        fixed_end = 10 if currency_split else 8
        navy = "123A5A"
        yellow = "FFF200"
        pale = "FFF2CC"
        amount_format = '#,##0.00;[Red](#,##0.00)'
        # Selected source amounts already feed the debit/credit values.  Keep
        # their individual columns in Source Details only; repeating them on
        # the printable voucher is misleading and makes the layout unreadable.
        table_end = fixed_end
        thin = Side(style="thin", color="808080")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        widths = (
            {"A": 20, "B": 32, "C": 18, "D": 18, "E": 38,
             "F": 20, "G": 32, "H": 18, "I": 18, "J": 38}
            if currency_split else
            {"A": 20, "B": 32, "C": 18, "D": 38,
             "E": 20, "F": 32, "G": 18, "H": 38}
        )
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        if configure_sheet:
            sheet.sheet_view.showGridLines = False
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.page_setup.orientation = "landscape"
            sheet.print_options.horizontalCentered = True

        summary = [
            ("Report ID", config.report_id_filter),
            ("Product Type", config.product_type),
            ("Reporting For", config.reporting_for),
            ("Rollup To", config.rollup_to or ""),
            ("Funds Xfer Entity", config.funds_xfer_entity or ""),
            ("Production Date", fields.Date.to_string(config.production_date) if config.production_date else ""),
            ("Report Date", fields.Date.to_string(config.report_date) if config.report_date else ""),
            ("Settlement Currency", config.settlement_currency),
            ("Clearing Currency", config.clearing_currency or ""),
            ("Business Mode", config.business_mode or ""),
            ("Transaction Category", config.transaction_category or ""),
            ("Transaction Type", config.transaction_type or ""),
            ("VSS-120 Source Level", dict(config._fields["vss_120_source_level"].selection).get(config.vss_120_source_level, "")),
            ("Disposition", config.source_description or ""),
            ("Description", config.description_filter or ""),
            ("Particulars", config.particulars or ""),
            ("Entity", config.entity or ""),
        ]
        if currency_split:
            run = config.batch_run_id
            types = list(dict.fromkeys(
                config.line_ids.filtered(
                    lambda line: line.include_in_voucher
                    and (line.source_currency or "").strip().upper() == "USD"
                ).mapped("conversion_type")
            ))
            labels = dict(config.line_ids._fields["conversion_type"].selection)
            summary.extend([
                ("Conversion Types", ", ".join(labels.get(value, value) for value in types)),
                ("Currency Conversion GL", run.currency_conversion_gl_number if run else ""),
                ("USD Buy Rate", run.usd_buy_rate if run else 0.0),
                ("USD Sale Rate", run.usd_sale_rate if run else 0.0),
                ("USD Mid Rate", run.usd_mid_rate if run else 0.0),
            ])
        row = start_row
        for label, value in summary:
            label_cell = sheet.cell(row, 1, label)
            value_cell = sheet.cell(row, 2, value)
            label_cell.font = Font(bold=True, color=navy)
            label_cell.fill = PatternFill("solid", fgColor=pale)
            value_cell.fill = PatternFill("solid", fgColor=pale)
            label_cell.border = border
            value_cell.border = border
            label_cell.alignment = Alignment(vertical="center")
            value_cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = _voucher_row_height((value, 38))
            row += 1

        row += 1
        for title_index, title in enumerate((
            config.company_id.name,
            config.business_owner,
            config.reporting_for,
            f"Settlement (VSS) Date: {fields.Date.to_string(config.report_date)}",
            f"Posting Date: {fields.Date.to_string(config.posting_date)}",
        )):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
            cell = sheet.cell(row, 1, title)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(bold=title_index < 3, color=navy)
            sheet.row_dimensions[row].height = _voucher_row_height((title, 100))
            row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
        product_cell = sheet.cell(row, 1, config.product_type)
        product_cell.fill = PatternFill("solid", fgColor=yellow)
        product_cell.font = Font(bold=True)
        product_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 23
        row += 2

        header_row = row
        headers = (
            ["GL Number", "GL Name", "Debit Amount (USD)", "Debit Amount (BDT)", "Debit Remarks",
             "GL Number", "GL Name", "Credit Amount (USD)", "Credit Amount (BDT)", "Credit Remarks"]
            if currency_split else
            ["GL Number", "GL Name", "Debit Amount", "Debit Remarks",
             "GL Number", "GL Name", "Credit Amount", "Credit Remarks"]
        )
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        sheet.row_dimensions[row].height = 34
        row += 1

        debit_lines = voucher.line_ids.filtered(lambda line: line.posting_indicator == "DR").sorted("sequence")
        credit_lines = voucher.line_ids.filtered(lambda line: line.posting_indicator == "CR").sorted("sequence")
        max_lines = max(len(debit_lines), len(credit_lines), 1)
        for index in range(max_lines):
            if index < len(debit_lines):
                line = debit_lines[index]
                if currency_split:
                    values = (line.gl_number, line.gl_name, line.amount_usd, line.amount_bdt)
                    for offset, value in enumerate(values):
                        sheet.cell(row + index, 1 + offset, value)
                    sheet.cell(row + index, 1).number_format = "@"
                    sheet.cell(row + index, 3).number_format = amount_format
                    sheet.cell(row + index, 4).number_format = amount_format
                    sheet.cell(row + index, 5, line.debit_remarks or line.remarks or "")
                else:
                    values = (line.gl_number, line.gl_name, line.amount)
                    for offset, value in enumerate(values):
                        sheet.cell(row + index, 1 + offset, value)
                    sheet.cell(row + index, 1).number_format = "@"
                    sheet.cell(row + index, 3).number_format = amount_format
                    sheet.cell(row + index, 4, line.debit_remarks or line.remarks or "")
            if index < len(credit_lines):
                line = credit_lines[index]
                if currency_split:
                    values = (line.gl_number, line.gl_name, line.amount_usd, line.amount_bdt)
                    for offset, value in enumerate(values):
                        sheet.cell(row + index, 6 + offset, value)
                    sheet.cell(row + index, 6).number_format = "@"
                    sheet.cell(row + index, 8).number_format = amount_format
                    sheet.cell(row + index, 9).number_format = amount_format
                    sheet.cell(row + index, 10, line.credit_remarks or line.remarks or "")
                else:
                    values = (line.gl_number, line.gl_name, line.amount)
                    for offset, value in enumerate(values):
                        sheet.cell(row + index, 5 + offset, value)
                    sheet.cell(row + index, 5).number_format = "@"
                    sheet.cell(row + index, 7).number_format = amount_format
                    sheet.cell(row + index, 8, line.credit_remarks or line.remarks or "")
            for column in range(1, table_end + 1):
                cell = sheet.cell(row + index, column)
                cell.border = border
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="right" if column in ((3, 4, 8, 9) if currency_split else (3, 7)) else "left",
                    wrap_text=column in ((2, 5, 7, 10) if currency_split else (2, 4, 6, 8)),
                )
                if index % 2:
                    cell.fill = PatternFill("solid", fgColor="F7FAFC")
            debit_line = debit_lines[index] if index < len(debit_lines) else False
            credit_line = credit_lines[index] if index < len(credit_lines) else False
            sheet.row_dimensions[row + index].height = _voucher_row_height(
                (debit_line.gl_name if debit_line else "", 32),
                ((debit_line.debit_remarks or debit_line.remarks) if debit_line else "", 38),
                (credit_line.gl_name if credit_line else "", 32),
                ((credit_line.credit_remarks or credit_line.remarks) if credit_line else "", 38),
            )
        row += max_lines
        if currency_split:
            debit_usd = sum(debit_lines.mapped("amount_usd"))
            credit_usd = sum(credit_lines.mapped("amount_usd"))
            sheet.cell(row, 2, "Total").font = Font(bold=True)
            sheet.cell(row, 3, debit_usd).font = Font(bold=True)
            sheet.cell(row, 4, voucher.debit_total).font = Font(bold=True)
            sheet.cell(row, 7, "Total").font = Font(bold=True)
            sheet.cell(row, 8, credit_usd).font = Font(bold=True)
            sheet.cell(row, 9, voucher.credit_total).font = Font(bold=True)
            for column in (3, 4, 8, 9):
                sheet.cell(row, column).number_format = amount_format
        else:
            sheet.cell(row, 2, "Total").font = Font(bold=True)
            sheet.cell(row, 3, voucher.debit_total).font = Font(bold=True)
            sheet.cell(row, 6, "Total").font = Font(bold=True)
            sheet.cell(row, 7, voucher.credit_total).font = Font(bold=True)
            sheet.cell(row, 3).number_format = amount_format
            sheet.cell(row, 7).number_format = amount_format
        for column in range(1, table_end + 1):
            cell = sheet.cell(row, column)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
            cell.alignment = Alignment(
                horizontal="right" if column in ((3, 4, 8, 9) if currency_split else (3, 7)) else "left",
                vertical="center",
            )
        sheet.row_dimensions[row].height = 23
        row += 1
        if currency_split:
            diff_label_col, status_col = 7, 10
            differences = (
                (8, abs(debit_usd - credit_usd)),
                (9, abs(voucher.debit_total - voucher.credit_total)),
            )
        else:
            diff_label_col, status_col = 6, 8
            differences = ((7, voucher.difference),)
        sheet.cell(row, diff_label_col, "Difference").font = Font(bold=True)
        color = "008000" if abs(voucher.difference) < 0.005 else "C00000"
        for column, difference in differences:
            sheet.cell(row, column, difference)
            sheet.cell(row, column).number_format = amount_format
            sheet.cell(row, column).font = Font(bold=True, color=color)
        sheet.cell(row, status_col, "BALANCED" if abs(voucher.difference) < 0.005 else "MISMATCH FLAG")
        sheet.cell(row, status_col).font = Font(bold=True, color=color)
        sheet.cell(row, status_col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.row_dimensions[row].height = 26
        row += 2

        remarks = [config.remarks] + voucher.line_ids.mapped("remarks")
        remarks = list(dict.fromkeys(value.strip() for value in remarks if value and value.strip()))
        if remarks:
            remark_text = "Remarks: " + " | ".join(remarks)
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=table_end)
            sheet.cell(row, 1, remark_text)
            sheet.cell(row, 1).fill = PatternFill("solid", fgColor=pale)
            sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = _voucher_row_height((remark_text, 120))
            row += 3
        else:
            row += 2
        sheet.cell(row, 1, "Maker").font = Font(bold=True)
        sheet.cell(row, 2, self.generated_by_id.name or "")
        checker_label_col, checker_value_col = ((7, 8) if currency_split else (6, 7))
        sheet.cell(row, checker_label_col, "Checker").font = Font(bold=True)
        sheet.cell(row, checker_value_col, self.checked_by_id.name or "")
        if configure_sheet:
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.print_area = f"A{start_row}:{get_column_letter(table_end)}{row}"
        return row

    def _write_source_details(self, sheet):
        """Show selected amount sources as columns without changing posting lines."""
        config = self.configuration_id
        navy = "123A5A"
        groups, amount_fields = _group_source_details(self.voucher_id.line_ids)
        headers = [
            "Report ID", "Source Level", "Business Mode", "Transaction Category",
            "Transaction Type", "Source Description", "Source Currency", "Source CR/DB", "Posting DR/CR",
        ]
        headers.extend(AMOUNT_SOURCE_LABELS.get(field, field) for field in amount_fields or ("matched_amount",))
        headers.extend([
            "GL Number", "GL Name", "Source Row Count", "Parsed Row IDs",
            "Debit Remarks", "Credit Remarks",
        ])
        detail_end_letter = get_column_letter(len(headers))
        sheet.merge_cells(f"A1:{detail_end_letter}1")
        title = sheet.cell(1, 1, f"Voucher source details — {self.name}")
        title.font = Font(bold=True, size=14, color="FFFFFF")
        title.fill = PatternFill("solid", fgColor=navy)
        sheet.row_dimensions[1].height = 27
        sheet.merge_cells(f"A2:{detail_end_letter}2")
        sheet.cell(2, 1, f"{config.product_type}  |  {config.reporting_for}  |  {fields.Date.to_string(config.report_date)}")
        sheet.cell(2, 1).font = Font(color=navy, italic=True)
        for column, label in enumerate(headers, 1):
            cell = sheet.cell(4, column, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[4].height = 32
        widths = [15, 14, 16, 23, 24, 31, 17, 16, 17]
        widths.extend([19] * (len(amount_fields) or 1))
        widths.extend([17, 30, 17, 28, 28, 28])
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(4, column).column_letter].width = width
        amount_start = 10
        for row_number, group in enumerate(groups, 5):
            line = group["line"]
            values = [
                line.source_report_code, line.source_record_type, line.source_business_mode,
                line.source_transaction_category, line.source_transaction_type,
                line.description, line.source_currency,
                " / ".join(sorted(group["source_directions"])),
                " / ".join(sorted(group["posting_indicators"])),
            ]
            values.extend(group["amounts"].get(field, 0.0) for field in amount_fields or ("matched_amount",))
            values.extend([
                " / ".join(sorted(group["gl_numbers"])),
                " / ".join(sorted(group["gl_names"])),
                group["row_count"],
                ",".join(dict.fromkeys(group["record_ids"])),
                " | ".join(sorted(group["debit_remarks"])),
                " | ".join(sorted(group["credit_remarks"])),
            ])
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column, value)
                cell.alignment = Alignment(vertical="center")
                if row_number % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F7FA")
            for column in range(amount_start, amount_start + (len(amount_fields) or 1)):
                sheet.cell(row_number, column).number_format = '#,##0.00;[Red](#,##0.00)'
            sheet.row_dimensions[row_number].height = 20
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:{detail_end_letter}{max(sheet.max_row, 5)}"
        sheet.sheet_view.showGridLines = False

    def action_download(self):
        self.ensure_one()
        if not self.output_file:
            self._generate_output_file()
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{self._name}/{self.id}/output_file/{self.output_filename}?download=true",
            "target": "self",
        }

    def action_regenerate_file(self):
        self.action_generate()
        return self.action_download()

    def action_mark_checked(self):
        for report in self:
            if report.state != "balanced":
                raise UserError(_("Only a balanced report can be marked as checked."))
            if report.generated_by_id == self.env.user:
                raise ValidationError(_("Maker and Checker must be different users."))
            report.write({"state": "checked", "checked_by_id": self.env.user.id, "checked_at": fields.Datetime.now()})
            report._generate_output_file()

    def action_approve(self):
        for report in self:
            if report.state != "checked":
                raise UserError(_("The report must be checked before approval."))
            report.write({"state": "approved", "approved_by_id": self.env.user.id, "approved_at": fields.Datetime.now()})

    def action_cancel(self):
        self.write({"state": "cancelled"})


class BsrAccountingVoucher(models.Model):
    _name = "bsr.accounting.voucher"
    _description = "Generated Voucher Block"
    _order = "report_id desc, id"
    _check_company_auto = True

    report_id = fields.Many2one("bsr.accounting.report", required=True, index=True, ondelete="cascade")
    company_id = fields.Many2one(related="report_id.company_id", store=True, index=True)
    product_type = fields.Char(related="report_id.product_type", store=True)
    line_ids = fields.One2many("bsr.accounting.voucher.line", "voucher_id", copy=False)
    issue_ids = fields.One2many("bsr.accounting.voucher.issue", "voucher_id", copy=False)
    state = fields.Selection([("balanced", "Balanced"), ("exception", "Exception")], default="exception", required=True)
    debit_total = fields.Float(compute="_compute_totals", store=True, digits=(20, 2))
    credit_total = fields.Float(compute="_compute_totals", store=True, digits=(20, 2))
    difference = fields.Float(compute="_compute_totals", store=True, digits=(20, 2))
    usd_total = fields.Float(compute="_compute_totals", store=True, digits=(20, 2))
    line_count = fields.Integer(compute="_compute_totals", store=True)
    issue_count = fields.Integer(compute="_compute_totals", store=True)
    show_currency_columns = fields.Boolean(related="report_id.configuration_id.show_currency_columns", store=True)

    @api.depends("line_ids.amount", "line_ids.amount_bdt", "line_ids.amount_usd", "line_ids.posting_indicator", "issue_ids")
    def _compute_totals(self):
        for voucher in self:
            debit_lines = voucher.line_ids.filtered(lambda line: line.posting_indicator == "DR")
            credit_lines = voucher.line_ids.filtered(lambda line: line.posting_indicator == "CR")
            voucher.debit_total = sum(debit_lines.mapped("amount"))
            voucher.credit_total = sum(credit_lines.mapped("amount"))
            bdt_difference = abs(voucher.debit_total - voucher.credit_total)
            usd_difference = abs(
                sum(debit_lines.mapped("amount_usd"))
                - sum(credit_lines.mapped("amount_usd"))
            )
            voucher.difference = (
                max(usd_difference, bdt_difference)
                if voucher.show_currency_columns
                else bdt_difference
            )
            voucher.usd_total = sum(voucher.line_ids.mapped("usd_amount"))
            voucher.line_count = len(voucher.line_ids)
            voucher.issue_count = len(voucher.issue_ids)

    def _build_from_configuration(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.issue_ids.unlink()
        config = self.report_id.configuration_id
        split_currency = bool(config.show_currency_columns)
        run = config.batch_run_id
        rules = config.line_ids.filtered("include_in_voucher").sorted(lambda item: (item.sequence, item.id))
        values_list = []
        for rule in rules:
            if not rule.matched_amount:
                self.env["bsr.accounting.voucher.issue"].create(
                    {
                        "voucher_id": self.id,
                        "mapping_id": rule.id,
                        "severity": "warning",
                        "message": _("The configured source rule has no matched amount."),
                    }
                )
                continue
            amount_usd, amount_bdt, conversion_rate, remaining_amounts = _mapping_currency_values(
                rule, run
            )
            requires_conversion = bool(amount_usd and not _paired_amount_field(rule))
            if requires_conversion and conversion_rate <= 0:
                raise UserError(
                    _("A positive %(type)s is required for USD rule %(rule)s.")
                    % {
                        "type": dict(rule._fields["conversion_type"].selection).get(
                            rule.conversion_type, _("conversion rate")
                        ),
                        "rule": rule.display_name,
                    }
                )
            values_list.append(
                {
                    "voucher_id": self.id,
                    "sequence": rule.sequence,
                    "mapping_id": rule.id,
                    "gl_number": rule.gl_number,
                    "gl_name": rule.gl_name,
                    "posting_indicator": rule.posting_indicator,
                    "amount": amount_bdt,
                    "amount_usd": amount_usd,
                    "amount_bdt": amount_bdt,
                    "conversion_rate": conversion_rate,
                    "conversion_type": rule.conversion_type if requires_conversion else False,
                    "usd_amount": rule.usd_amount,
                    "description": rule.description_match or rule.transaction_type or rule.report_code,
                    "debit_remarks": rule.debit_remarks
                    or (rule.remarks if rule.posting_indicator == "DR" else False),
                    "credit_remarks": rule.credit_remarks
                    or (rule.remarks if rule.posting_indicator == "CR" else False),
                    "remarks": (
                        rule.debit_remarks
                        if rule.posting_indicator == "DR"
                        else rule.credit_remarks
                    )
                    or rule.remarks,
                    "source_report_code": rule.report_code,
                    "source_record_type": rule.record_type,
                    "source_business_mode": rule.business_mode,
                    "source_transaction_category": rule.transaction_category,
                    "source_transaction_type": rule.transaction_type,
                    "source_amount_field": rule.amount_source,
                    "amount_columns_json": remaining_amounts,
                    "source_currency": _mapping_amount_currency(rule, rule.amount_source),
                    "source_direction": rule.source_direction,
                    "source_row_count": rule.source_row_count,
                    "source_record_ids": rule.source_record_ids,
                }
            )
        if split_currency and values_list:
            debit_usd = sum(value["amount_usd"] for value in values_list if value["posting_indicator"] == "DR")
            credit_usd = sum(value["amount_usd"] for value in values_list if value["posting_indicator"] == "CR")
            debit_bdt = sum(value["amount_bdt"] for value in values_list if value["posting_indicator"] == "DR")
            credit_bdt = sum(value["amount_bdt"] for value in values_list if value["posting_indicator"] == "CR")
            adjustments = _currency_balance_adjustments(
                debit_usd, credit_usd, debit_bdt, credit_bdt
            )
            if adjustments:
                if not run or not run.currency_conversion_gl_number:
                    raise UserError(_("Enter the Currency Conversion GL before generating the voucher."))
            next_sequence = max((value["sequence"] for value in values_list), default=0)
            for adjustment in adjustments:
                next_sequence += 10
                posting_indicator = adjustment["side"]
                values_list.append(
                    {
                        "voucher_id": self.id,
                        "sequence": next_sequence,
                        "gl_number": run.currency_conversion_gl_number,
                        "gl_name": adjustment["label"],
                        "posting_indicator": posting_indicator,
                        "amount": adjustment["amount_bdt"],
                        "amount_usd": adjustment["amount_usd"],
                        "amount_bdt": adjustment["amount_bdt"],
                        "conversion_rate": 0.0,
                        "description": adjustment["label"],
                        "debit_remarks": adjustment["label"] if posting_indicator == "DR" else False,
                        "credit_remarks": adjustment["label"] if posting_indicator == "CR" else False,
                        "remarks": adjustment["label"],
                        "source_report_code": "FX-ADJUSTMENT",
                        "source_currency": "/".join(
                            currency
                            for currency, amount in (
                                ("USD", adjustment["amount_usd"]),
                                ("BDT", adjustment["amount_bdt"]),
                            )
                            if amount >= 0.005
                        ),
                        "is_conversion_adjustment": True,
                    }
                )
        if values_list:
            self.env["bsr.accounting.voucher.line"].create(values_list)
        self.state = "balanced" if self.line_ids and not self.issue_ids and abs(self.difference) < 0.005 else "exception"


class BsrAccountingVoucherLine(models.Model):
    _name = "bsr.accounting.voucher.line"
    _description = "Generated Voucher Line"
    _order = "voucher_id desc, sequence, id"

    voucher_id = fields.Many2one("bsr.accounting.voucher", required=True, index=True, ondelete="cascade")
    show_currency_columns = fields.Boolean(related="voucher_id.show_currency_columns", store=True)
    company_id = fields.Many2one(related="voucher_id.company_id", store=True, index=True)
    sequence = fields.Integer(default=10)
    mapping_id = fields.Many2one("bsr.accounting.mapping", ondelete="set null", index=True)
    gl_number = fields.Char(required=True)
    gl_name = fields.Char(required=True)
    posting_indicator = fields.Selection([("DR", "DR"), ("CR", "CR")], required=True, index=True)
    amount = fields.Float(required=True, digits=(20, 2))
    amount_usd = fields.Float(string="Amount (USD)", digits=(20, 2))
    amount_bdt = fields.Float(string="Amount (BDT)", digits=(20, 2))
    conversion_rate = fields.Float(string="USD BDT Rate", digits=(20, 6))
    conversion_type = fields.Selection(
        [("buy", "Buy Rate"), ("sale", "Sale Rate"), ("mid", "Mid Rate")],
        string="Conversion Type",
    )
    is_conversion_adjustment = fields.Boolean(string="Currency Gain/Loss", readonly=True)
    usd_amount = fields.Float(string="USD Amount", digits=(20, 2))
    description = fields.Char()
    remarks = fields.Text()
    debit_remarks = fields.Text(string="Debit Remarks")
    credit_remarks = fields.Text(string="Credit Remarks")
    source_report_code = fields.Char()
    source_record_type = fields.Char()
    source_business_mode = fields.Char()
    source_transaction_category = fields.Char()
    source_transaction_type = fields.Char()
    source_amount_field = fields.Char()
    amount_columns_json = fields.Json(
        string="Additional Amount Columns",
        readonly=True,
        copy=True,
        help="Informational source amounts selected after the primary accounting amount.",
    )
    source_currency = fields.Char()
    source_direction = fields.Char()
    source_row_count = fields.Integer()
    source_record_ids = fields.Text()


class BsrAccountingVoucherIssue(models.Model):
    _name = "bsr.accounting.voucher.issue"
    _description = "Voucher Generation Issue"
    _order = "voucher_id desc, id"

    voucher_id = fields.Many2one("bsr.accounting.voucher", required=True, index=True, ondelete="cascade")
    company_id = fields.Many2one(related="voucher_id.company_id", store=True, index=True)
    mapping_id = fields.Many2one("bsr.accounting.mapping", ondelete="set null", index=True)
    severity = fields.Selection([("warning", "Warning"), ("error", "Error")], default="warning", required=True)
    message = fields.Text(required=True)

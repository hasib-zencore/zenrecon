# -*- coding: utf-8 -*-
"""Parsers for configurable CMS/CBS reconciliation sources.

The staging layer deliberately stores a stable JSON representation.  It accepts
normal tabular bank extracts and VISA VX text reports, which are useful as
realistic test fixtures for CMS/CBS during implementation and UAT.
"""

import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from .vx705_parser import parse_vx705
from .vx706_parser import parse_vx706
from .vx707_parser import parse_vx707
from .vx725_parser import parse_vx725
from .vx727_parser import parse_vx727


VX_PARSERS = {
    "705": parse_vx705,
    "706": parse_vx706,
    "707": parse_vx707,
    "725": parse_vx725,
    "727": parse_vx727,
}


def column_key(label, used=None):
    """Return a predictable technical key while preserving duplicate headers."""
    key = re.sub(r"[^a-z0-9]+", "_", str(label or "").strip().lower()).strip("_")
    key = key or "column"
    if used is None:
        return key
    base = key
    suffix = 2
    while key in used:
        key = "%s_%s" % (base, suffix)
        suffix += 1
    used.add(key)
    return key


def _json_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (str, int, float, bool)):
        return str(value).strip()
    return str(value)


def _headers(labels):
    used = set()
    return [
        {"name": str(label or "Column %s" % (index + 1)).strip(), "key": column_key(label, used)}
        for index, label in enumerate(labels)
    ]


def _rows_from_matrix(matrix):
    nonempty = [row for row in matrix if any(_json_value(cell) for cell in row)]
    if len(nonempty) < 2:
        raise ValueError("The file must contain a header row and at least one data row.")
    headers = _headers(nonempty[0])
    rows = []
    for row_number, cells in enumerate(nonempty[1:], 2):
        values = {
            header["key"]: _json_value(cells[index] if index < len(cells) else "")
            for index, header in enumerate(headers)
        }
        if any(values.values()):
            rows.append({"row_number": row_number, "values": values})
    if not rows:
        raise ValueError("No data rows were found in the uploaded file.")
    return headers, rows, "tabular"


def _parse_xlsx(raw):
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    return _rows_from_matrix([list(row) for row in sheet.iter_rows(values_only=True)])


def _parse_json(raw):
    payload = json.loads(raw.decode("utf-8-sig"))
    if isinstance(payload, dict):
        for key in ("records", "rows", "data", "items"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list) or not payload or not all(isinstance(row, dict) for row in payload):
        raise ValueError("JSON input must be an object or a list of objects.")
    labels = []
    for row in payload:
        for label in row:
            if label not in labels:
                labels.append(label)
    headers = _headers(labels)
    label_to_key = {label: header["key"] for label, header in zip(labels, headers)}
    rows = [
        {
            "row_number": index,
            "values": {label_to_key[label]: _json_value(row.get(label)) for label in labels},
        }
        for index, row in enumerate(payload, 1)
    ]
    return headers, rows, "json"


def _decode_text(raw):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_vx(text, filename):
    match = re.search(r"\bREPORT\s+VX-(705|706|707|725|727)\b", text, re.IGNORECASE)
    if not match:
        return None
    report_type = match.group(1)
    values, _unmatched = VX_PARSERS[report_type](text, source_file_name=filename)
    if not values:
        raise ValueError("No transaction rows were found in the VX-%s file." % report_type)
    keys = []
    for row in values:
        for key in row:
            if key not in keys and key not in {"company_id", "batch_id"}:
                keys.append(key)
    headers = [{"name": key.replace("_", " ").title(), "key": key} for key in keys]
    rows = []
    for index, row in enumerate(values, 1):
        rows.append({
            "row_number": int(row.get("source_line_number") or index),
            "values": {key: _json_value(row.get(key)) for key in keys},
        })
    return headers, rows, "vx%s" % report_type


def _parse_delimited(text):
    sample = "\n".join(text.splitlines()[:25])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in sample else csv.excel
    matrix = list(csv.reader(io.StringIO(text), dialect))
    return _rows_from_matrix(matrix)


def parse_dynamic_file(raw, filename=""):
    """Return ``(headers, rows, format_name)`` for a supported upload."""
    if not raw:
        raise ValueError("The uploaded file is empty.")
    extension = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if extension in {"xlsx", "xlsm"}:
        return _parse_xlsx(raw)
    if extension == "json":
        return _parse_json(raw)
    text = _decode_text(raw)
    vx_result = _parse_vx(text, filename)
    if vx_result:
        return vx_result
    return _parse_delimited(text)
